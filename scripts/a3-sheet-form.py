# Z曲げ・ハット曲げ・コの字（中押しの要否）の実績記入シートを A3 横1枚で作る。
#   計算値は scripts/a3-sheet-sim.mjs の出力（JSON）を使う。
#   使い方: python scripts/a3-sheet-form.py <Z|HAT|UNAKA> <sim.json> <出力.xlsx> <図.png>
import json
import sys
from datetime import date

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import rcParams
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

kind, sim_path, out_path, fig_path = sys.argv[1].upper(), sys.argv[2], sys.argv[3], sys.argv[4]
rows = json.load(open(sim_path, encoding='utf-8'))
FONT = 'Meiryo UI'
TODAY = date.today().isoformat()

# ---------------------------------------------------------------- 図（寸法の取り方）
rcParams['font.family'] = ['Meiryo', 'Yu Gothic', 'MS Gothic']


def arrow(ax, x1, y1, x2, y2, txt, tx, ty):
    ax.annotate('', xy=(x1, y1), xytext=(x2, y2), arrowprops=dict(arrowstyle='<->', color='#555', lw=1.2))
    ax.text(tx, ty, txt, ha='center', va='center', fontsize=12, fontweight='bold')


if kind == 'Z':
    fig, axs = plt.subplots(1, 2, figsize=(9.2, 2.3), gridspec_kw={'width_ratios': [1.3, 1]})
    a = axs[0]
    a.plot([0, 60, 60, 130], [60, 60, 0, 0], color='#111', lw=5, solid_joinstyle='miter')
    arrow(a, 0, 74, 60, 74, 'フランジA', 30, 86)
    arrow(a, 76, 0, 76, 60, '段差S', 100, 30)
    arrow(a, 60, -16, 130, -16, 'フランジB', 95, -28)
    a.text(105, 30, 'すべて外寸', fontsize=10, color='#444', ha='center')
    a.set_xlim(-14, 150); a.set_ylim(-40, 100)
    b = axs[1]
    b.plot([0, 60, 60, 130], [60, 60, 0, 0], color='#bbb', lw=5, solid_joinstyle='miter')
    b.plot([0, 60], [60, 60], color='#c00', lw=5)
    b.text(30, 74, '1曲げ目', color='#c00', fontsize=11, ha='center', fontweight='bold')
    b.text(95, -22, '2曲げ目（裏返し）', color='#444', fontsize=11, ha='center')
    b.set_xlim(-14, 150); b.set_ylim(-40, 100)
    title = 'Z曲げ 実績記入シート'
    intro = ('書き方　① 寸法はすべて外寸（左の絵）。A と B は同じ長さで見ます。'
             '② 「記入」の列に、実際に曲げられた値を書いてください。③ 曲げられなかったときは、当たった所（ヤゲン・中間板・ホルダ・柱・ダイ）を備考に。')
    ent = [('段差S 最小\n（実際）', 13), ('そのSでの\nフランジA上限', 13), ('フランジA 最大\n（段差S=100）', 13), ('確かめた人', 11), ('日付', 10), ('備考（当たった所 など）', 26)]
    sims = [('段差S 最小', 'sMin'), ('フランジA上限\nS=30', 'a30'), ('フランジA上限\nS=50', 'a50'), ('フランジA上限\nS=100', 'a100')]
elif kind == 'HAT':
    fig, axs = plt.subplots(1, 2, figsize=(9.2, 2.3), gridspec_kw={'width_ratios': [1.5, 1]})
    a = axs[0]
    a.plot([0, 40, 40, 110, 110, 150], [0, 0, 60, 60, 0, 0], color='#111', lw=5, solid_joinstyle='miter')
    arrow(a, 0, -16, 40, -16, '外フランジA', 20, -30)
    arrow(a, 110, -16, 150, -16, '外フランジB', 130, -30)
    arrow(a, 26, 0, 26, 60, '立上りH1', 8, 30)
    arrow(a, 124, 0, 124, 60, '立上りH2', 142, 30)
    arrow(a, 40, 74, 110, 74, '上面W', 75, 86)
    a.text(75, 30, 'すべて外寸。左右同じ寸法で見ます', fontsize=9, color='#444', ha='center')
    a.set_xlim(-20, 170); a.set_ylim(-42, 100)
    b = axs[1]
    b.plot([0, 40, 40, 110, 110, 150], [0, 0, 60, 60, 0, 0], color='#bbb', lw=5, solid_joinstyle='miter')
    b.plot([40, 40], [0, 60], color='#c00', lw=5)
    b.plot([110, 110], [0, 60], color='#c00', lw=5)
    b.text(75, 74, '内側2か所を先に、外側2か所をあとで', color='#c00', fontsize=10, ha='center', fontweight='bold')
    b.set_xlim(-20, 170); b.set_ylim(-42, 100)
    title = 'ハット曲げ 実績記入シート'
    intro = ('書き方　① 寸法はすべて外寸（左の絵）。左右同じ寸法（A=B、H1=H2）で見ます。外フランジは 50mm で計算しています。'
             '② ハットは立上りが低すぎても（上面が上型に近づいて）当たるので、曲げられる立上りは「下限〜上限」の範囲で出しています。'
             '③ 実際に曲げられた値を「記入」に。曲げられなかったときは当たった所を備考に。')
    ent = [('上面W 最小\n（実際）', 13), ('立上りH 下限\n（W=100）', 13), ('立上りH 上限\n（W=100）', 13), ('確かめた人', 11), ('日付', 10), ('備考（外フランジ・当たった所）', 24)]
    sims = [('上面W 最小\n（H=50）', 'wMin'), ('立上りH の範囲\nW=50', 'h50'), ('立上りH の範囲\nW=100', 'h100'), ('立上りH の範囲\nW=200', 'h200')]
else:
    fig, axs = plt.subplots(1, 3, figsize=(9.6, 2.3))
    a = axs[0]
    a.plot([0, 0, 90, 90], [70, 0, 0, 70], color='#111', lw=5, solid_joinstyle='miter')
    arrow(a, -16, 0, -16, 70, '立上りH', -40, 35)
    arrow(a, 0, -16, 90, -16, '底W', 45, -30)
    a.text(45, 35, 'すべて外寸\n左右同じ高さ', fontsize=9, color='#444', ha='center')
    a.set_xlim(-60, 110); a.set_ylim(-42, 96)
    b = axs[1]
    b.plot([0, 0, 90, 90], [70, 0, 0, 70], color='#2a6fbf', lw=5, solid_joinstyle='miter')
    b.text(45, 86, '普通に曲げる', color='#2a6fbf', fontsize=12, ha='center', fontweight='bold')
    b.text(45, 35, '立上りが高いと\n上型に当たる', fontsize=9, color='#444', ha='center')
    b.set_xlim(-20, 110); b.set_ylim(-20, 104)
    c = axs[2]
    c.plot([0, 0, 45, 90, 90], [70, 0, 10, 0, 70], color='#6f4bc4', lw=5, solid_joinstyle='miter')
    c.annotate('', xy=(45, 12), xytext=(45, 50), arrowprops=dict(arrowstyle='->', color='#6f4bc4', lw=2))
    c.text(45, 86, '中押し（捨て曲げ）', color='#6f4bc4', fontsize=12, ha='center', fontweight='bold')
    c.text(45, 70, 'への字 → 両サイド → 押し戻す', fontsize=9, color='#444', ha='center')
    c.set_xlim(-20, 110); c.set_ylim(-20, 104)
    title = 'コの字曲げ　中押しが要る範囲 確認シート'
    intro = ('見方　立上りを高くしていくと ① そのまま曲がる → ② 中押し（捨て曲げ）が要る → ③ どちらでも曲がらない、の順になります。'
             '「中押しが要る範囲」のいちばん下の数字から試して、中押しなしで曲がったら「記入」に書いてください。'
             '内-内が120mm未満の中押しは現場で未確認です。')
    ent = [('中押しなしで曲がった\n立上りの最大（W=100）', 16), ('中押しが要った\n立上りの最小（W=100）', 16),
           ('中押しでも無理だった\n立上り（W=100）', 16), ('確かめた人', 11), ('日付', 10), ('備考（底W・当たった所）', 20)]
    sims = [('そのまま曲がる\nW=50', 'n50'), ('中押しが要る範囲\nW=50', 'z50'),
            ('そのまま曲がる\nW=100', 'n100'), ('中押しが要る範囲\nW=100', 'z100'),
            ('そのまま曲がる\nW=150', 'n150'), ('中押しが要る範囲\nW=150', 'z150'),
            ('そのまま曲がる\nW=200', 'n200'), ('中押しが要る範囲\nW=200', 'z200')]

for ax in axs:
    ax.set_aspect('equal')
    ax.axis('off')
fig.tight_layout(pad=0.2)
fig.savefig(fig_path, dpi=200, facecolor='white')
plt.close(fig)

# ---------------------------------------------------------------- 表
wb = Workbook()
ws = wb.active
ws.title = {'Z': 'Z曲げ実績', 'HAT': 'ハット曲げ実績', 'UNAKA': 'コの字 中押し'}[kind]

thin = Side(style='thin', color='999999')
med = Side(style='medium', color='444444')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='E8EBE6')
ENTRY = PatternFill('solid', fgColor='FFF7E6')     # 記入欄＝うすい黄
SIMF = PatternFill('solid', fgColor='EAF1FB')      # シミュレーション＝うすい青
MATF = PatternFill('solid', fgColor='F0F0EC')

base_cols = [('下型', 8), ('V幅', 6), ('曲げ内R', 7), ('材質', 6), ('板厚 t', 7), ('最小フランジ\n外寸（表）', 11), ('機械', 10)]
cols = base_cols + [(h, w) for h, w in ent] + [(h, 11) for h, _ in sims]
for i, (_, w) in enumerate(cols, start=1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
last = len(cols)


def put(r, c, v, *, bold=False, size=10, fill=None, wrap=True, align='center'):
    cell = ws.cell(r, c, v)
    cell.font = Font(name=FONT, size=size, bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    cell.border = box
    if fill:
        cell.fill = fill
    return cell


# 見出し
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
t = ws.cell(1, 1, f'{title}　―　株式会社高橋鉄骨')
t.font = Font(name=FONT, size=16, bold=True)
t.alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[1].height = 26

ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
n = ws.cell(2, 1, f'{intro}　（シミュレーションの値は {TODAY} 時点・ヤゲン904061・中間板標準。うすい青＝参考、うすい黄＝記入欄）')
n.font = Font(name=FONT, size=9)
n.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws.row_dimensions[2].height = 26

# 図（3〜12行目）
img = XLImage(fig_path)
img.width, img.height = (760, 190) if kind != 'UNAKA' else (800, 190)
ws.add_image(img, 'A3')
for r in range(3, 13):
    ws.row_dimensions[r].height = 15

# 表の見出し（14行目＝大見出し、15行目＝小見出し）
HEAD_ROW, SUB_ROW = 14, 15
ws.merge_cells(start_row=HEAD_ROW, start_column=1, end_row=HEAD_ROW, end_column=len(base_cols))
put(HEAD_ROW, 1, '金型・材質・板厚', bold=True, fill=HEAD, align='left')
c0 = len(base_cols) + 1
ws.merge_cells(start_row=HEAD_ROW, start_column=c0, end_row=HEAD_ROW, end_column=c0 + len(ent) - 1)
put(HEAD_ROW, c0, '記 入 ── 実際の値（曲げ屋さんに確かめて書く）', bold=True, fill=ENTRY, align='left')
c1 = c0 + len(ent)
ws.merge_cells(start_row=HEAD_ROW, start_column=c1, end_row=HEAD_ROW, end_column=last)
put(HEAD_ROW, c1, 'シミュレーション（参考）', bold=True, fill=SIMF, align='left')
ws.row_dimensions[HEAD_ROW].height = 20

for i, (h, _) in enumerate(cols, start=1):
    fill = HEAD if i <= len(base_cols) else ENTRY if i < c1 else SIMF
    put(SUB_ROW, i, h, bold=True, size=9, fill=fill)
ws.row_dimensions[SUB_ROW].height = 44

# 中押しが要る範囲＝「そのまま曲がる上限」を超えてから「中押しの上限」まで
def zone(row, W):
    n, m = row.get(f'n{W}'), row.get(f'm{W}')
    lo = n if isinstance(n, (int, float)) else None
    if isinstance(m, str) and m != '上限なし':
        return '不可'
    hi = m
    if lo is None:
        return f'〜{hi}' if hi is not None else '—'
    if hi is None:
        return '—'
    if hi == '上限なし':
        return f'{round(lo) + 1}〜'
    if float(hi) <= float(lo):
        return '利点なし'
    return f'{round(lo) + 1}〜{round(float(hi))}'


if kind == 'UNAKA':
    for row in rows:
        for W in (50, 100, 150, 200):
            row[f'z{W}'] = zone(row, W)

# 行
r = SUB_ROW + 1
cur_mat = None
for row in rows:
    if row['mat'] != cur_mat:
        cur_mat = row['mat']
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
        put(r, 1, f'材質：{cur_mat}', bold=True, fill=MATF, align='left')
        ws.row_dimensions[r].height = 18
        r += 1
    put(r, 1, f"V{row['V']}", bold=True)
    put(r, 2, row['V'])
    put(r, 3, row.get('r'))
    put(r, 4, row['mat'])
    put(r, 5, row['t'])
    put(r, 6, row['minOut'])
    put(r, 7, row.get('machine'), size=9)
    for i in range(len(ent)):
        put(r, len(base_cols) + 1 + i, None, fill=ENTRY)
    for i, (_, key) in enumerate(sims):
        v = row.get(key)
        put(r, c1 + i, v if v is not None else '—', size=9, fill=SIMF)
    ws.row_dimensions[r].height = 17
    r += 1

# 下の余白（品物ごとの実例）
r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
put(r, 1, '品物ごとの実例（曲がった／曲がらなかった、どちらも書いてください）', bold=True, fill=HEAD, align='left')
r += 1
# 見出しごとに何列分を使うか（合計が表の列数になるように最後で調整）
ex = ([('日付', 1), ('品名・図番', 2), ('材質', 1), ('板厚', 1), ('下型(V)', 1), ('立上りH・底W（外寸）', 2),
       ('曲げ長さ L', 1), ('中押しなし ○／✕', 2), ('中押し ○／✕', 1), ('当たった所・ひとこと', None)]
      if kind == 'UNAKA' else
      [('日付', 1), ('品名・図番', 2), ('材質', 1), ('板厚', 1), ('下型(V)', 1), ('寸法（外寸）', 2),
       ('曲げ長さ L', 1), ('曲げ方（普通／くの字／中押し）', 2), ('結果（○／✕）', 1), ('当たった所・ひとこと', None)])
spans = []
used = 0
for name, sp in ex:
    n = (last - used) if sp is None else sp
    spans.append((name, max(1, n)))
    used += max(1, n)
c = 1
for name, sp in spans:
    if c > last:
        break
    end = min(last, c + sp - 1)
    if end > c:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=end)
    put(r, c, name, bold=True, size=9, fill=HEAD)
    c = end + 1
ws.row_dimensions[r].height = 26
for rr in range(r + 1, r + 9):
    c = 1
    for _, sp in spans:
        if c > last:
            break
        end = min(last, c + sp - 1)
        if end > c:
            ws.merge_cells(start_row=rr, start_column=c, end_row=rr, end_column=end)
        for cc in range(c, end + 1):
            put(rr, cc, None, fill=ENTRY)
        c = end + 1
    ws.row_dimensions[rr].height = 20

# 印刷設定：A3横・1ページに収める
ws.page_setup.paperSize = 8          # A3
ws.page_setup.orientation = 'landscape'
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.35, bottom=0.3, header=0.2, footer=0.2)
ws.print_title_rows = f'{HEAD_ROW}:{SUB_ROW}'
ws.freeze_panes = ws.cell(SUB_ROW + 1, 1)
wb.save(out_path)
print('saved', out_path)
