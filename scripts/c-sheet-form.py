# C形曲げ（リップ内向き）の実績入力シートを作る。1ファイル2シート、どちらも A3 横1枚。
#   1枚目「型ごと」  ：型 × 材質 × 板厚ごとに、決めた底W・立上りHで曲げられたリップの範囲（最小〜最大）を書く。
#                       計算値は scripts/c-sheet-sim.mjs の出力（JSON）を参考に並べる。
#   2枚目「品物ごと」：曲げた品物を1件1行で書く（そのまま実績 bendsim.json の zuRecords に登録できる項目）。
#   使い方: node scripts/c-sheet-sim.mjs > c.json
#           python scripts/c-sheet-form.py c.json <出力.xlsx> <図.png>
import json
import sys
from datetime import date

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import rcParams
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

sim_path, out_path, fig_path = sys.argv[1], sys.argv[2], sys.argv[3]
rows = json.load(open(sim_path, encoding='utf-8'))
# 材質（鉄 → 縞）・V・板厚の順に並べる（JSON は 5・6 が 3.2 より先に来ることがある）
rows.sort(key=lambda r: (['鉄', '縞'].index(r['mat']), r['V'], r['t']))
FONT = 'Meiryo UI'
TODAY = date.today().isoformat()
CASES = [(100, 50), (150, 75), (200, 100)]   # c-sheet-sim.mjs と同じ

# ---------------------------------------------------------------- 図（寸法の取り方）
rcParams['font.family'] = ['Meiryo', 'Yu Gothic', 'MS Gothic']


def arrow(ax, x1, y1, x2, y2):
    ax.annotate('', xy=(x1, y1), xytext=(x2, y2), arrowprops=dict(arrowstyle='<->', color='#555', lw=1.2))


fig, axs = plt.subplots(2, 1, figsize=(3.4, 4.6))
a = axs[0]
a.plot([30, 0, 0, 110, 110, 80], [70, 70, 0, 0, 70, 70], color='#111', lw=5, solid_joinstyle='miter')
arrow(a, 0, 84, 30, 84); a.text(15, 96, 'リップA', ha='center', fontsize=11, fontweight='bold')
arrow(a, 80, 84, 110, 84); a.text(95, 96, 'リップB', ha='center', fontsize=11, fontweight='bold')
arrow(a, -14, 0, -14, 70); a.text(-30, 35, '立上り\nH1', ha='center', va='center', fontsize=11, fontweight='bold')
arrow(a, 124, 0, 124, 70); a.text(141, 35, '立上り\nH2', ha='center', va='center', fontsize=11, fontweight='bold')
arrow(a, 0, -14, 110, -14); a.text(55, -27, '底W', ha='center', va='center', fontsize=11, fontweight='bold')
a.text(55, 32, 'すべて外寸', fontsize=10, color='#444', ha='center')
a.set_xlim(-48, 158); a.set_ylim(-38, 106)
b = axs[1]
b.plot([30, 0, 0, 110, 110, 80], [70, 70, 0, 0, 70, 70], color='#bbb', lw=5, solid_joinstyle='miter')
b.plot([30, 0], [70, 70], color='#c00', lw=5)
b.plot([110, 80], [70, 70], color='#c00', lw=5)
b.text(55, 88, 'リップを長くすると、最後の曲げで\nヤゲンに当たる・抜けなくなる', color='#c00', fontsize=9.5, ha='center', fontweight='bold')
b.annotate('', xy=(55, 66), xytext=(55, 20), arrowprops=dict(arrowstyle='->', color='#777', lw=1.5))
b.text(55, 8, 'ヤゲン', color='#777', fontsize=9, ha='center')
b.set_xlim(-48, 158); b.set_ylim(-10, 110)
for ax in axs:
    ax.set_aspect('equal')
    ax.axis('off')
fig.tight_layout(pad=0.2)
fig.savefig(fig_path, dpi=200, facecolor='white')
plt.close(fig)

# ---------------------------------------------------------------- 共通の部品
thin = Side(style='thin', color='999999')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='E8EBE6')
ENTRY = PatternFill('solid', fgColor='FFF7E6')     # 記入欄＝うすい黄
SIMF = PatternFill('solid', fgColor='EAF1FB')      # シミュレーション＝うすい青
MATF = PatternFill('solid', fgColor='F0F0EC')
EXF = PatternFill('solid', fgColor='F4F4F4')       # 記入例＝うすい灰


def put(ws, r, c, v, *, bold=False, size=13, fill=None, wrap=True, align='center', color=None, italic=False):
    cell = ws.cell(r, c, v)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    cell.border = box
    if fill:
        cell.fill = fill
    return cell


def widths(ws, cols, target):
    # A3横（余白込み）に収まる文字数の目安。合計が target になるよう、列幅をまとめて伸ばす
    scale = target / sum(w for _, w in cols)
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(4.5, w * scale)


def header(ws, title, intro, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    t = ws.cell(1, 1, f'{title}　―　株式会社高橋鉄骨')
    t.font = Font(name=FONT, size=17, bold=True)
    t.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 23
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    n = ws.cell(2, 1, intro)
    n.font = Font(name=FONT, size=9.5)
    n.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 46


def a3(ws, head_rows):
    ws.page_setup.paperSize = 8          # A3
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.35, bottom=0.3, header=0.2, footer=0.2)
    ws.print_title_rows = head_rows


def figure(ws, col):
    img = XLImage(fig_path)
    img.width, img.height = 200, 270
    ws.column_dimensions[get_column_letter(col - 1)].width = 2
    ws.column_dimensions[get_column_letter(col)].width = 29
    ws.add_image(img, f'{get_column_letter(col)}4')


wb = Workbook()

# ---------------------------------------------------------------- 1枚目：型ごと
ws = wb.active
ws.title = '型ごと'
base_cols = [('下型', 10), ('材質', 8), ('板厚 t', 9), ('最小フランジ\n外寸（表）', 13), ('機械', 17)]
ent = [(f'リップ 最小〜最大\nW={W}・H={H}', 19) for W, H in CASES] + [('備考（当たった所・\n抜けなかった など）', 30)]
sims = [(f'リップの範囲\nW={W}・H={H}', f'a{W}') for W, H in CASES]
cols = [(h, w * 1.5) for h, w in base_cols] + [(h, w * 1.5) for h, w in ent] + [(h, 19) for h, _ in sims]
widths(ws, cols, 158)
last = len(cols)
header(ws, 'C形曲げ 実績入力シート（型ごと）',
       '書き方　① 寸法はすべて外寸（右の絵）。左右同じ寸法（A=B、H1=H2）で見ます。'
       '② 「記入」の列に、その底W・立上りHで、曲げられたリップの長さを「最小〜最大」で書いてください（例 20〜45。曲げられなかったら「不可」）。リップは長すぎても短すぎても当たります。'
       '③ 曲げられなかったときは、当たった所（ヤゲン・中間板・ホルダ・柱・ダイ）や「曲げた後にヤゲンから抜けなかった」を備考に。'
       f'　（シミュレーションの値は {TODAY} 時点・ヤゲン904061・中間板標準。V12〜V25 はZ曲げで実際より甘く出た型。うすい青＝参考、うすい黄＝記入欄）',
       last + 2)
ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(6, last // 2))
who = ws.cell(3, 1, '確かめた人：　　　　　　　　　　　　　　　日付：　　　　年　　月　　日')
who.font = Font(name=FONT, size=11, bold=True)
who.alignment = Alignment(horizontal='left', vertical='center')
who.border = Border(bottom=thin)
ws.row_dimensions[3].height = 19
figure(ws, last + 2)

HEAD_ROW, SUB_ROW = 4, 5
ws.merge_cells(start_row=HEAD_ROW, start_column=1, end_row=HEAD_ROW, end_column=len(base_cols))
put(ws, HEAD_ROW, 1, '金型・材質・板厚', bold=True, fill=HEAD, align='left')
c0 = len(base_cols) + 1
ws.merge_cells(start_row=HEAD_ROW, start_column=c0, end_row=HEAD_ROW, end_column=c0 + len(ent) - 1)
put(ws, HEAD_ROW, c0, '記 入 ── 実際の値（曲げ屋さんに確かめて書く）', bold=True, fill=ENTRY, align='left')
c1 = c0 + len(ent)
ws.merge_cells(start_row=HEAD_ROW, start_column=c1, end_row=HEAD_ROW, end_column=last)
put(ws, HEAD_ROW, c1, 'シミュレーション（参考）', bold=True, fill=SIMF, align='left')
ws.row_dimensions[HEAD_ROW].height = 17
for i, (h, _) in enumerate(cols, start=1):
    put(ws, SUB_ROW, i, h, bold=True, size=10.5, fill=HEAD if i < c0 else ENTRY if i < c1 else SIMF)
ws.row_dimensions[SUB_ROW].height = 50

r = SUB_ROW + 1
cur_mat = None
for row in rows:
    if row['mat'] != cur_mat:
        cur_mat = row['mat']
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
        put(ws, r, 1, f'材質：{cur_mat}', bold=True, fill=MATF, align='left')
        ws.row_dimensions[r].height = 18
        r += 1
    put(ws, r, 1, f"V{row['V']}", bold=True)
    put(ws, r, 2, row['mat'])
    put(ws, r, 3, row['t'])
    put(ws, r, 4, row['minOut'])
    put(ws, r, 5, row.get('machine'), size=11)
    for i in range(len(ent)):
        put(ws, r, c0 + i, None, fill=ENTRY)
    for i, (_, key) in enumerate(sims):
        v = row.get(key)
        put(ws, r, c1 + i, v if v is not None else '—', size=12.5, fill=SIMF)
    ws.row_dimensions[r].height = 16
    r += 1
a3(ws, f'{HEAD_ROW}:{SUB_ROW}')
ws.freeze_panes = ws.cell(SUB_ROW + 1, 1)

# ---------------------------------------------------------------- 2枚目：品物ごと
ws2 = wb.create_sheet('品物ごと')
cols2 = [('No', 4), ('日付', 9), ('伝票番号', 11), ('材質', 6), ('板厚 t', 6), ('下型', 9),
         ('リップA', 7), ('立上りH1', 7), ('底W', 7), ('立上りH2', 7), ('リップB', 7), ('曲げ長さ L', 8),
         ('結果\n○／✕', 7), ('曲げ方', 10), ('✕のときの理由', 14), ('当たった所・ひとこと', 24), ('確かめた人', 10)]
last2 = len(cols2)
widths(ws2, cols2, 158)
header(ws2, 'C形曲げ 実績入力シート（品物ごと）',
       '書き方　曲げた品物を1件1行で書いてください。① 寸法はすべて外寸（右の絵）。② 結果は 曲がった＝○、曲がらなかった＝✕。'
       '③ 材質・下型・結果・曲げ方・理由は、セルを選ぶと出る ▼ から選べます（手書きでもかまいません）。'
       '④ 曲がらなかったときは、理由（型に当たった／曲げた後にヤゲンから抜けなかった など）と当たった所を書いてください。'
       '1行目は記入例です（消して使ってください）。',
       last2 + 2)
figure(ws2, last2 + 2)
HR2 = 4
ws2.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
c = ws2.cell(3, 1, '寸法はすべて外寸（mm）')
c.font = Font(name=FONT, size=11, bold=True)
ws2.row_dimensions[3].height = 19
for i, (h, _) in enumerate(cols2, start=1):
    put(ws2, HR2, i, h, bold=True, size=10.5, fill=HEAD)
ws2.row_dimensions[HR2].height = 36

example = ['例', '10/10', '12345', '鉄', 3.2, 'V20', 20, 70, 100, 70, 20, 1970, '○', '普通', '', '記入例（消して使う）', '曲げ 田中']
N_ROWS = 25
for k in range(N_ROWS + 1):
    rr = HR2 + 1 + k
    for i in range(1, last2 + 1):
        if k == 0:
            put(ws2, rr, i, example[i - 1], size=11, fill=EXF, color='777777', italic=True,
                align='left' if i == 16 else 'center')
        else:
            put(ws2, rr, i, k if i == 1 else None, size=12, fill=None if i == 1 else ENTRY,
                align='left' if i == 16 else 'center')
    ws2.row_dimensions[rr].height = 27 if k else 22

# 選べる欄（▼）。リストにない書き方でも入れられるよう、エラーにはしない
dies = ['V8', 'V12', 'V12 2溝', 'V16', 'V16 2溝', 'V20', 'V20 2溝', 'V25', 'V25 2溝',
        'V32', 'V40', 'V50', 'V63', 'V80', 'V100', 'V125', 'V160']
first, end = HR2 + 2, HR2 + 1 + N_ROWS
for col, items in [(4, ['鉄', '縞']), (6, dies), (13, ['○', '✕']), (14, ['普通', 'くの字165', 'くの字100', '中押し']),
                   (15, ['型に当たった', '曲げた後 抜けない', '長さ（力）が足りない', 'その他'])]:
    dv = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True, showErrorMessage=False)
    ws2.add_data_validation(dv)
    L = get_column_letter(col)
    dv.add(f'{L}{first}:{L}{end}')
a3(ws2, f'{HR2}:{HR2}')
ws2.freeze_panes = ws2.cell(HR2 + 1, 1)

wb.save(out_path)
print('saved', out_path)
