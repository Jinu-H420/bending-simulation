# くの字ヤゲン 実績記入シートを A3 縦1枚で作る（曲げ屋さんに確かめて書いてもらう用）。
#   ① ヤゲンの寸法（図面の値と実物）  ② 型・板厚ごとの、くの字で曲げられる範囲  ③ 品物ごとの実例
#   計算値は scripts/kuno-sheet-sim.mjs の出力（JSON）を使う。
#   使い方: python scripts/kuno-sheet-form.py <sim.json> <出力.xlsx> <図.png>
import json
import sys
from datetime import date

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import rcParams
from matplotlib.patches import Polygon
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

sim_path, out_path, fig_path = sys.argv[1], sys.argv[2], sys.argv[3]
rows = json.load(open(sim_path, encoding='utf-8'))
R_OF = {8: 1.3, 12: 2, 16: 2.6, 20: 3.3, 25: 4, 32: 5, 40: 6.5, 50: 8, 63: 10, 80: 13, 100: 16, 125: 20, 160: 26}
FONT = 'Meiryo UI'

# ---------------------------------------------------------------- 図（くの字ヤゲンの正面と、曲げたときの横から）
rcParams['font.family'] = ['Meiryo', 'Yu Gothic', 'MS Gothic']
fig, axs = plt.subplots(1, 2, figsize=(11.5, 2.7), gridspec_kw={'width_ratios': [1.7, 1]})
for ax in axs:
    ax.set_aspect('equal'); ax.axis('off')
a = axs[0]
# 正面：全長330、両端65は全高97、中央200（窓）は高さ37。刃先が下（y=0）
E, WIN, H, M = 65, 200, 97, 37
pts = [(0, 0), (E + WIN + E, 0), (E + WIN + E, H), (E + WIN, H), (E + WIN, M), (E, M), (E, H), (0, H)]
a.add_patch(Polygon(pts, closed=True, facecolor='#dce6f1', edgecolor='#2a78d6', lw=2))
def dim(ax, x1, y1, x2, y2, txt, tx, ty, fs=11):
    ax.annotate('', xy=(x1, y1), xytext=(x2, y2), arrowprops=dict(arrowstyle='<->', color='#555', lw=1.1))
    ax.text(tx, ty, txt, ha='center', va='center', fontsize=fs, fontweight='bold')
dim(a, 0, -14, E, -14, '両端 65', E / 2, -26)
dim(a, E, -14, E + WIN, -14, '窓の長さ 200（100版は70）', E + WIN / 2, -26)
dim(a, E + WIN, -14, E + WIN + E, -14, '65', E + WIN + E / 2, -26)
dim(a, -14, 0, -14, H, '97', -30, H / 2)
dim(a, E + 18, 0, E + 18, M, '窓の高さ 37', E + 60, M / 2 + 2, fs=10)
a.text(E + WIN / 2, M + 22, 'ここに品物を入れる（L は窓の長さまで）', ha='center', fontsize=10, color='#c62f2f')
a.set_xlim(-45, 345); a.set_ylim(-38, 115)
a.set_title('くの字ヤゲン（正面）　全長 330（100版は200）', fontsize=11)
b = axs[1]
# 横から：先に立てた立上りが、低い窓の上を通る
b.add_patch(Polygon([(-6, 0), (6, 0), (10, 37), (-10, 37)], closed=True, facecolor='#dce6f1', edgecolor='#2a78d6', lw=2))
b.plot([-60, -2, 2, 60], [60, 2, 2, 60], color='#111', lw=4, solid_joinstyle='miter')
b.plot([-60, -100], [60, 20], color='#111', lw=4)
b.text(0, 50, '立上りが\n窓の上を通る', ha='center', fontsize=9.5, color='#444')
b.set_xlim(-110, 80); b.set_ylim(-10, 110)
b.set_title('横から（2か所目を曲げるとき）', fontsize=11)
plt.tight_layout()
plt.savefig(fig_path, dpi=160)
plt.close()

# ---------------------------------------------------------------- シート
wb = Workbook()
ws = wb.active
ws.title = 'くの字ヤゲン実績'
thin = Side(style='thin', color='999999')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def fill(h): return PatternFill('solid', fgColor=h)
C_HEAD, C_KEY, C_REF, C_CALC, C_IN, C_SEC, C_KUNO = 'BFBFBF', 'EAEAEA', 'DCE6F1', 'E2EFDA', 'FFF2CC', 'D9D9D9', 'DDEBF7'
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
widths = {'A': 7, 'B': 5, 'C': 5.5, 'D': 5, 'E': 5.5, 'F': 8,
          'G': 8, 'H': 9, 'I': 8, 'J': 9, 'K': 8, 'L': 9, 'M': 8, 'N': 9,
          'O': 9, 'P': 10, 'Q': 26}
for k, v in widths.items():
    ws.column_dimensions[k].width = v
LAST = 'Q'

def put(cell, v, bg=None, bold=False, size=10, al=center, color='000000', italic=False):
    c = ws[cell]
    c.value = v
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    c.alignment = al
    if bg: c.fill = fill(bg)
    return c
def box(rng, bg=None):
    if ':' not in rng: rng = f'{rng}:{rng}'
    for row in ws[rng]:
        for c in row:
            c.border = BOX
            if bg: c.fill = fill(bg)
def merge(rng, v, **kw):
    ws.merge_cells(rng)
    put(rng.split(':')[0], v, **kw)
    box(rng, kw.get('bg'))

# 見出し
ws.row_dimensions[1].height = 28
ws.merge_cells(f'A1:{LAST}1')
put('A1', 'くの字ヤゲン 実績記入シート　―　特殊 くの字165・くの字100', size=16, bold=True, al=left)
ws.row_dimensions[2].height = 18
ws.merge_cells(f'A2:{LAST}2')
put('A2', f'株式会社高橋鉄骨　曲げ可否判断シート用　／　作成 {date.today():%Y-%m-%d}　／　'
          '色：灰＝型と板（さわらない）　青＝図面の値　緑＝計算の値（さわらない）　黄＝実際の値を書く欄', size=10, al=left)

# 図と書き方
for r in range(3, 10):
    ws.row_dimensions[r].height = 22.5
img = XLImage(fig_path)
img.width, img.height = 650, 153
ws.add_image(img, 'A3')
merge(f'N3:{LAST}9',
      '書き方\n'
      '① 緑・青＝計算や図面の値。その右の黄の欄に、実際の値を書く。分かる所だけでよい。\n'
      '② 計算どおりなら「〃」、上限が無いなら「なし」。\n'
      '③ くの字を使わない型・板厚は「使う？」に ✕。\n'
      '④ いちばん大事なのは 曲げ長さ L。実際に曲げられた一番長い L を書く。\n'
      '⑤ 品物ごとの記録は下の「実例」へ。',
      bg='F2F6FB', size=9.5, al=Alignment(horizontal='left', vertical='top', wrap_text=True))

# ① ヤゲンの寸法
r = 11
ws.row_dimensions[r].height = 18
merge(f'A{r}:{LAST}{r}', '① ヤゲンの寸法（図面の値と実物）', bg=C_HEAD, bold=True, al=left)
r += 1
ws.row_dimensions[r].height = 30
heads1 = [('A', 'C', 'ヤゲン'), ('D', 'E', '全長\n図面'), ('F', 'F', '全長\n実際'), ('G', 'G', '窓の長さ\n図面'),
          ('H', 'H', '窓の長さ\n実際'), ('I', 'I', '窓の高さ\n図面'), ('J', 'J', '窓の高さ\n実際'),
          ('K', 'K', '両端\n図面'), ('L', 'L', '両端\n実際'), ('M', 'M', '本数\n実際'), ('N', 'O', '使える機械\n（HG・HD）'),
          ('P', 'P', '曲げられる\n最大板厚'), ('Q', 'Q', '備考')]
for a_, b_, name in heads1:
    bg = C_REF if '図面' in name else C_IN if name != 'ヤゲン' else C_KEY
    merge(f'{a_}{r}:{b_}{r}', name, bg=bg, bold=True, size=8.5)
r += 1
for name, total, win in [('特殊 くの字165', 330, 200), ('特殊 くの字100', 200, 70)]:
    ws.row_dimensions[r].height = 22
    vals = {'A': name, 'D': total, 'G': win, 'I': 37, 'K': 65}
    for a_, b_, h in heads1:
        if a_ != b_: ws.merge_cells(f'{a_}{r}:{b_}{r}')
        v = vals.get(a_)
        bg = C_KEY if a_ == 'A' else C_REF if a_ in vals else C_IN
        put(f'{a_}{r}', v, bg=bg, bold=(a_ == 'G'), al=left if a_ in 'AQ' else center)
        box(f'{a_}{r}:{b_}{r}', bg)
    r += 1

# ② 型・板厚ごとの範囲
r += 1
ws.row_dimensions[r].height = 18
merge(f'A{r}:{LAST}{r}', '② 型・板厚ごとに、くの字で曲げられる範囲（コの字・2か所とも90°。計算は くの字165・左右の立上りは同じ高さ）', bg=C_HEAD, bold=True, al=left)
r += 1
H2, H3 = r, r + 1
ws.row_dimensions[H2].height = 36
ws.row_dimensions[H3].height = 16
for col, name in zip('ABCDE', ['下型', 'V幅', '曲げ\n内R', '材質', '板厚\nt']):
    merge(f'{col}{H2}:{col}{H3}', name, bg=C_KEY, bold=True, size=9)
merge(f'F{H2}:F{H3}', 'くの字を\n使う？\n○／✕', bg=C_IN, bold=True, size=8.5)
pairs = [('G', '底W 最小\n（立上り最短）'), ('I', '立上りH 上限\n底W＝50'), ('K', '立上りH 上限\n底W＝100'), ('M', '立上りH 上限\n底W＝200')]
for col, name in pairs:
    col2 = chr(ord(col) + 1)
    merge(f'{col}{H2}:{col2}{H2}', name, bg=C_KUNO, bold=True, size=9)
    put(f'{col}{H3}', '計算', bg=C_CALC, bold=True, size=8.5); box(f'{col}{H3}')
    put(f'{col2}{H3}', '実際', bg=C_IN, bold=True, size=8.5); box(f'{col2}{H3}')
merge(f'O{H2}:O{H3}', '曲げられた\n最大の L\n（実際）', bg=C_IN, bold=True, size=8.5)
merge(f'P{H2}:P{H3}', '普通の\n904061なら\n(底W100)', bg=C_REF, bold=True, size=8)
merge(f'Q{H2}:Q{H3}', '備考\n（当たった所など）', bg=C_IN, bold=True, size=9)
r = H3 + 1
first_data = r
# 普通のヤゲンの立上り上限（底W=100）：くの字と比べるため。コの字実績記入シート2と同じ計算
try:
    base = {(x['V'], x['mat'], x['t']): x for x in json.load(open(sys.argv[4], encoding='utf-8'))} if len(sys.argv) > 4 else {}
except Exception:
    base = {}
def section(label):
    global r
    ws.row_dimensions[r].height = 15
    merge(f'A{r}:{LAST}{r}', label, bg=C_SEC, bold=True, size=9.5, al=left)
    r += 1
fx = lambda v: '不可' if v == 'W不足' else v
for mat in ['鉄', '縞']:
    section(f'材質：{mat}')
    for d in [x for x in rows if x['mat'] == mat]:
        ws.row_dimensions[r].height = 20
        V, t = d['V'], d['t']
        for col, v in [('A', f'V{V}'), ('B', V), ('C', R_OF.get(V)), ('D', mat), ('E', t)]:
            put(f'{col}{r}', v, bg=C_KEY); box(f'{col}{r}')
        put(f'F{r}', None, bg=C_IN); box(f'F{r}')
        for col, v in {'G': d['wMin'], 'I': fx(d['h50']), 'K': fx(d['h100']), 'M': fx(d['h200'])}.items():
            put(f'{col}{r}', v, bg=C_CALC); box(f'{col}{r}')
            put(f'{chr(ord(col) + 1)}{r}', None, bg=C_IN); box(f'{chr(ord(col) + 1)}{r}')
        put(f'O{r}', None, bg=C_IN); box(f'O{r}')
        b = base.get((V, mat, t))
        put(f'P{r}', fx(b['h100']) if b else None, bg=C_REF, size=9, color='555555'); box(f'P{r}')
        put(f'Q{r}', None, bg=C_IN, al=left, size=9); box(f'Q{r}')
        r += 1
section('その他（表に無い材質・板厚）')
for _ in range(2):
    ws.row_dimensions[r].height = 20
    for col in 'ABCDE': put(f'{col}{r}', None, bg=C_KEY); box(f'{col}{r}')
    for col in 'FHJLNOQ': put(f'{col}{r}', None, bg=C_IN); box(f'{col}{r}')
    for col in 'GIKMP': put(f'{col}{r}', None, bg=C_CALC if col != 'P' else C_REF); box(f'{col}{r}')
    r += 1
last_data = r - 1
orange = PatternFill('solid', fgColor='F8CBAD')
for col in 'GIKM':
    col2 = chr(ord(col) + 1)
    ws.conditional_formatting.add(f'{col2}{first_data}:{col2}{last_data}', FormulaRule(
        formula=[f'AND(ISNUMBER({col2}{first_data}),ISNUMBER({col}{first_data}),ABS({col2}{first_data}-{col}{first_data})>=3)'], fill=orange))

# ③ 実例
r += 1
ws.row_dimensions[r].height = 18
merge(f'A{r}:{LAST}{r}', '③ くの字で曲げた実例（品物ごと。曲がった・曲がらなかった、どちらも書く）', bg=C_HEAD, bold=True, al=left)
r += 1
ws.row_dimensions[r].height = 30
heads3 = [('A', 'B', '日付'), ('C', 'D', '材質'), ('E', 'E', '板厚'), ('F', 'F', '下型'), ('G', 'H', 'ヤゲン\n165／100'),
          ('I', 'I', '底W'), ('J', 'J', 'H1'), ('K', 'K', 'H2'), ('L', 'M', '曲げ長さ L'), ('N', 'N', '結果\n○／✕'),
          ('O', 'P', 'だめなとき\n当たった所'), ('Q', 'Q', 'メモ')]
for a_, b_, name in heads3:
    merge(f'{a_}{r}:{b_}{r}', name, bg=C_KEY, bold=True, size=9)
r += 1
ex_first = r
for _ in range(7):
    ws.row_dimensions[r].height = 20
    for a_, b_, _ in heads3:
        if a_ != b_: ws.merge_cells(f'{a_}{r}:{b_}{r}')
        box(f'{a_}{r}:{b_}{r}', C_IN)
        ws[f'{a_}{r}'].alignment = center
    r += 1
ex = ['2026/9/22', 'ボンデ', 1.2, 'V12', '165', 48, 121, 121, 100, '○', '', '記入例（消して使う）']
for (a_, _, _), v in zip(heads3, ex):
    put(f'{a_}{ex_first}', v, bg=C_IN, italic=True, color='808080', size=9)

# 注記
r += 1
notes = [
    '※ 緑の計算の値は、曲げシミュレーター（くの字165＝00402 の断面を高さ37mmで切った形・V.dxf で確認した台）の干渉判定です。実測ではありません。',
    '※ くの字は品物を中央の窓に入れて曲げるので、曲げ長さ L は窓の長さ（165版 200mm・100版 70mm）までとして計算しています。ここがいちばん確かめたい所です。',
    '※ 「不可」＝その底Wでは最短の立上りでも曲げられない。青の「普通の904061なら」は比べるための値（コの字実績記入シート2と同じ）。',
    '※ 現場の順番：普通に曲げる → くの字（L が窓に入るとき）→ 中押しは最終手段。',
    '※ Excelで入力したとき、実際の値が計算と3mm以上ちがうと、その欄がオレンジになります。',
]
for n in notes:
    ws.row_dimensions[r].height = 15
    ws.merge_cells(f'A{r}:{LAST}{r}')
    put(f'A{r}', n, size=9, al=left)
    r += 1

ws.print_area = f'A1:{LAST}{r - 1}'
ws.print_title_rows = f'{H2}:{H3}'
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.orientation = 'portrait'
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.35, bottom=0.35, header=0.15, footer=0.15)
ws.print_options.horizontalCentered = True
ws.freeze_panes = f'G{first_data}'
ws.sheet_view.zoomScale = 90
wb.save(out_path)
print('saved', out_path, 'rows', r - 1)
