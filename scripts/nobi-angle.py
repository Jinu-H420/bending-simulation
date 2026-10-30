# 角度ごとの片伸び（伸び一覧表）を Excel から読み、bending-simulator.jsx の NOBI_ANGLE に書き込む。
#
#   python scripts/nobi-angle.py            … 共有フォルダの表を読む
#   python scripts/nobi-angle.py 別の.xlsx  … ほかの表を読む
#
# 表：\\srv02\共有\データフォルダ\ユーザー共有用\地主\claudecode\曲げサンプルデータ\
#      折り曲げ_伸び一覧表_鋭角のみ入力枠分割版※入力.xlsx（シート「折り曲げ_伸び一覧表(印刷用)」）
#  ・「角度 / 板厚」の行が見出し。C列から右へ「3.2t (V20)」の形の列を読む（見出しが空の列で止める）
#  ・材質は見出しの1行上の題（「縞」が入っていれば縞、ほかは鉄）
#  ・鋭角（85°以下）は「仮想交点」の行だけ読む（シミュレーターの外寸は仮想交点まで）
#  ・90°の行の「-3.0」のような文字は、符号を取った値（片伸びは正で控除）
#  ・空欄は書かない（シミュレーターが前後の角度から割合で出す）
# Excel を開いたままでも読めるように、いったん写してから読む。
import os, re, shutil, sys, tempfile
import openpyxl

SRC = r'\\srv02\共有\データフォルダ\ユーザー共有用\地主\claudecode\曲げサンプルデータ\折り曲げ_伸び一覧表_鋭角のみ入力枠分割版※入力.xlsx'
JSX = os.path.join(os.path.dirname(__file__), '..', 'bending-simulator.jsx')
BEGIN, END = '// <NOBI_ANGLE 自動生成', '// </NOBI_ANGLE>'

src = sys.argv[1] if len(sys.argv) > 1 else SRC
tmp = os.path.join(tempfile.mkdtemp(), 'nobi.xlsx')
shutil.copy(src, tmp)
ws = openpyxl.load_workbook(tmp, data_only=True).worksheets[0]

def num(v):
    if v is None or (isinstance(v, str) and not v.strip()): return None
    try: return abs(float(str(v).strip()))
    except ValueError: return None

table, warns = {}, []   # table[材質][V][t] = {角度: 値}
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if not (isinstance(a, str) and a.startswith('角度')): continue
    title = str(ws.cell(r - 1, 1).value or '')
    mat = '縞' if '縞' in title else '鉄'
    cols = []
    for c in range(3, ws.max_column + 1):
        h = ws.cell(r, c).value
        if h is None or not str(h).strip(): break
        m = re.match(r'\s*([\d.]+)\s*t\s*\(\s*V\s*(\d+)\s*\)', str(h).replace('\n', ''))
        if m: cols.append((c, float(m.group(1)), int(m.group(2))))
    rr = r + 1
    while rr <= ws.max_row:
        lab = ws.cell(rr, 1).value
        if isinstance(lab, str) and lab.startswith('角度'): break
        m = re.match(r'\s*(\d+)\s*°', str(lab or ''))
        if m:
            ang = int(m.group(1))
            kind = str(ws.cell(rr, 2).value or '')
            if ang < 90 and '仮想' not in kind:
                rr += 1; continue
            for c, t, V in cols:
                v = num(ws.cell(rr, c).value)
                if v is not None:
                    table.setdefault(mat, {}).setdefault(V, {}).setdefault(t, {})[ang] = v
        rr += 1

# 気になる値（90°から鈍角へ行くと伸びは減るはずなのに、増えているところ）
for mat, byV in table.items():
    for V, byT in byV.items():
        for t, pts in byT.items():
            obt = sorted(a for a in pts if a >= 90)
            for a0, a1 in zip(obt, obt[1:]):
                if pts[a1] > pts[a0] + 0.05:
                    warns.append(f'{mat} t{t:g} V{V}：{a0}°={pts[a0]:g} → {a1}°={pts[a1]:g}（増えている）')

def js(x): return f'{x:g}'
out = [f'{BEGIN}（scripts/nobi-angle.py。手で直さない。表を直したら python scripts/nobi-angle.py）']
out.append('const NOBI_ANGLE = {')
for mat in sorted(table):
    out.append(f'  {mat}: {{')
    for V in sorted(table[mat]):
        out.append(f'    {V}: {{')
        for t in sorted(table[mat][V]):
            pts = table[mat][V][t]
            body = ', '.join(f'{a}: {js(pts[a])}' for a in sorted(pts))
            out.append(f'      {js(t)}: {{ {body} }},')
        out.append('    },')
    out.append('  },')
out.append('};')
out.append(f'const NOBI_ANGLE_WARN = {warns!r};'.replace("'", '"') if warns else 'const NOBI_ANGLE_WARN = [];')
out.append(END)

text = open(JSX, encoding='utf-8').read()
i, j = text.find(BEGIN), text.find(END)
if i < 0 or j < 0: sys.exit('bending-simulator.jsx に NOBI_ANGLE の印がありません')
j += len(END)
text = text[:i] + '\n'.join(out) + text[j:]
open(JSX, 'w', encoding='utf-8', newline='\n').write(text)

n = sum(len(byT) for byV in table.values() for byT in byV.values())
print(f'{n} 列を書き込みました（{", ".join(f"{m} {sum(len(b) for b in table[m].values())}列" for m in sorted(table))}）')
for w in warns: print('確認：', w)
