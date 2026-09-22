# 「小さいV 最長L 確認シート」を A3 縦1枚で作る（曲げ屋さんに書いてもらう用）。
#   板厚に対して基準（折り曲げ表の赤枠）より小さいVで曲げるとき、曲げられる最長の曲げ長さ L を書いてもらう。
#   行は scripts/smallv-sheet.mjs の出力。書いてもらった値は bending-simulator.jsx の SMALLV_MAXL に入れる。
#   使い方: python scripts/smallv-sheet-form.py <smallv.json> <出力.xlsx>
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

rows = json.load(open(sys.argv[1], encoding='utf-8'))
out_path = sys.argv[2]
FONT = 'Meiryo UI'
wb = Workbook()
ws = wb.active
ws.title = '小さいV 最長L'
thin = Side(style='thin', color='999999')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def fill(h): return PatternFill('solid', fgColor=h)
C_HEAD, C_KEY, C_REF, C_IN, C_SEC = 'BFBFBF', 'EAEAEA', 'DCE6F1', 'FFF2CC', 'D9D9D9'
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
cols = [('A', 6, '材質'), ('B', 7, '板厚 t'), ('C', 8, '使う V'), ('D', 8, '基準の V\n(折り曲げ表)'), ('E', 11, '機械'),
        ('F', 20, 'ダイ'), ('G', 10, '力の目安\n(t/m・計算)'), ('H', 11, '機械の力で\n決まる L\n(参考)'),
        ('I', 12, '曲げられた\n最長の L\n(実際)'), ('J', 12, '曲げられ\nなかった L\n(実際)'), ('K', 10, 'そのときの形\nZ／コ／L'),
        ('L', 30, '備考（金型がもたない・たわむ など）')]
for c, w, _ in cols:
    ws.column_dimensions[c].width = w
LAST = 'L'

def put(cell, v, bg=None, bold=False, size=11, al=center, color='000000', italic=False):
    c = ws[cell]
    c.value = v
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    c.alignment = al
    if bg: c.fill = fill(bg)
    c.border = BOX
    return c

ws.row_dimensions[1].height = 30
ws.merge_cells(f'A1:{LAST}1')
put('A1', '小さいV で曲げるときの 最長L 確認シート', size=16, bold=True, al=left).border = Border()
ws.row_dimensions[2].height = 18
ws.merge_cells(f'A2:{LAST}2')
put('A2', f'株式会社高橋鉄骨　曲げ可否判断シート用　／　作成 {date.today():%Y-%m-%d}　／　色：灰＝型と板　青＝計算の目安（参考）　黄＝実際の値を書く欄',
    size=10, al=left).border = Border()
ws.row_dimensions[3].height = 64
ws.merge_cells(f'A3:{LAST}3')
put('A3', '書き方\n'
          '① 板厚に対して基準より小さいVで曲げると、大きな力が要るので長いものは曲げられません。その「曲げられた一番長い L」を黄の欄に書いてください。\n'
          '② 曲げられなかった長さが分かれば、それも書いてください。分かる所だけでよい。使わない組み合わせは備考に「使わない」。\n'
          '③ 青の数字は機械の力（HG 220t・HD 350t）だけで計算した参考です。金型がもつかどうかは入っていません。',
    bg='F2F6FB', size=10, al=Alignment(horizontal='left', vertical='top', wrap_text=True))

r = 5
ws.row_dimensions[r].height = 44
for c, _, name in cols:
    bg = C_IN if c in 'IJKL' else C_REF if c in 'GH' else C_KEY
    put(f'{c}{r}', name, bg=bg, bold=True, size=9.5)
r += 1
first = r
for mat in ['鉄', '縞']:
    sub = [x for x in rows if x['mat'] == mat]
    if not sub: continue
    ws.row_dimensions[r].height = 18
    ws.merge_cells(f'A{r}:{LAST}{r}')
    put(f'A{r}', f'材質：{mat}', bg=C_SEC, bold=True, size=10, al=left)
    for c, _, _ in cols: ws[f'{c}{r}'].border = BOX
    r += 1
    for d in sub:
        ws.row_dimensions[r].height = 38
        vals = {'A': mat, 'B': d['t'], 'C': f"V{d['V']}", 'D': f"V{d['baseV']}", 'E': d['machine'], 'F': d['die'],
                'G': d['P'], 'H': d['Lcap'], 'I': d.get('maxL'), 'J': None, 'K': None, 'L': None}
        for c, _, _ in cols:
            bg = C_IN if c in 'IJKL' else C_REF if c in 'GH' else C_KEY
            put(f'{c}{r}', vals[c], bg=bg, bold=(c == 'C'), al=left if c in 'FL' else center, size=11 if c in 'BCDI' else 10)
        r += 1
ws.row_dimensions[r].height = 18
ws.merge_cells(f'A{r}:{LAST}{r}')
put(f'A{r}', 'その他（表に無い組み合わせ）', bg=C_SEC, bold=True, size=10, al=left)
for c, _, _ in cols: ws[f'{c}{r}'].border = BOX
r += 1
for _ in range(5):
    ws.row_dimensions[r].height = 38
    for c, _, _ in cols:
        put(f'{c}{r}', None, bg=C_IN if c in 'IJKL' else C_KEY)
    r += 1
# 記入例（その他の1行目）
for c, v in {'A': '鉄', 'B': 6, 'C': 'V20', 'D': 'V40', 'E': 'HD3504NT', 'F': '30540 2溝 V20溝', 'I': 1000, 'J': 1500, 'K': 'コ', 'L': '記入例（消して使う）'}.items():
    ws[f'{c}{r - 5}'].value = v
    ws[f'{c}{r - 5}'].font = Font(name=FONT, size=10, italic=True, color='808080')

r += 1
for n in ['※ 基準の V は折り曲げ表（ベンダー折り曲げ金型寸法表）の赤枠の型です。これより小さいVは、曲げに要る力が V に反比例して大きくなります。',
          '※ 力の目安は 1.42×400×板厚²÷V（kN/m）を t/m にしたもの（SS400）。参考なので、実際の値を優先します。',
          '※ 書いてもらった L は、シミュレーターと Z・コの字判定に入れ、その長さまでなら ○、超えたら ✕ にします。いまは「確認中」で △ にしています。']:
    ws.row_dimensions[r].height = 16
    ws.merge_cells(f'A{r}:{LAST}{r}')
    c = ws[f'A{r}']
    c.value = n
    c.font = Font(name=FONT, size=9.5)
    c.alignment = left
    r += 1

ws.print_area = f'A1:{LAST}{r - 1}'
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.orientation = 'portrait'
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
ws.print_options.horizontalCentered = True
ws.freeze_panes = f'A{first}'
wb.save(out_path)
print('saved', out_path, 'rows', r - 1)
