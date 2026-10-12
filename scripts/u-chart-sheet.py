# コの字曲げ 早見グラフ（Excel のグラフ機能で作る。A3横）
#   早見表（u-table-sheet.py）と同じ数字を、Excel のグラフにする。画像ではないので、点にマウスを当てると数字が出て、拡大もできる。
#   1枚目「全部の型」：16型の上限の線を重ねる。2枚目から 1型1枚：大きいグラフ（5mm ごとの点、10mm ごとに数字、
#   下がる所は赤い点）と、右に 5mm 刻みの表（立上りの下限〜上限）。縦軸は低いほうの立上り（左右違うときは低いほうで見る）。
#   使い方: node scripts/u-table-sim.mjs > ut.json
#           python scripts/u-chart-sheet.py ut.json <出力.xlsx>
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

sim_path, out_path = sys.argv[1], sys.argv[2]
rows = json.load(open(sim_path, encoding='utf-8'))
rows.sort(key=lambda r: (r['t'], r['base']))
FONT = 'Meiryo UI'
TODAY = date.today().isoformat()
H_TOP = 300

thin = Side(style='thin', color='999999')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='E8EBE6')
MINF = PatternFill('solid', fgColor='DDF2DD')
WF = PatternFill('solid', fgColor='F4F1E8')
COLORS = ['1F77B4', 'FF7F0E', '2CA02C', 'D62728', '9467BD', '8C564B', 'E377C2', '7F7F7F',
          'BCBD22', '17BECF', '393B79', 'AD494A', '637939', '8C6D31', '843C39', '7B4173']


def put(ws, r, c, v, *, bold=False, size=11, fill=None, color=None, align='center'):
    cell = ws.cell(r, c, v)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    cell.border = box
    if fill:
        cell.fill = fill
    return cell


def a3(ws, area):
    ws.page_setup.paperSize = 8
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.35, bottom=0.3, header=0.2, footer=0.2)
    ws.print_area = area


def title(ws, text, note, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    t = ws.cell(1, 1, text)
    t.font = Font(name=FONT, size=16, bold=True)
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    n = ws.cell(2, 1, note)
    n.font = Font(name=FONT, size=10)
    n.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 44


def line_props(color, width=2.25, dash=None, none=False):
    gp = GraphicalProperties()
    if none:
        gp.line.noFill = True
    else:
        gp.line.solidFill = color
        gp.line.width = int(width * 12700)
        if dash:
            gp.line.dashStyle = dash
    return gp


def axes(ch, xmin, xmax, ymax):
    ch.x_axis.title = '底W（外寸 mm）'
    ch.y_axis.title = '低いほうの立上り（外寸 mm）'
    ch.x_axis.scaling.min, ch.x_axis.scaling.max = xmin, xmax
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, ymax
    ch.x_axis.majorUnit = 50
    ch.y_axis.majorUnit = 50
    ch.x_axis.minorUnit = 10
    ch.y_axis.minorUnit = 10
    from openpyxl.chart.axis import ChartLines
    ch.x_axis.majorGridlines = ChartLines()
    ch.x_axis.minorGridlines = ChartLines()
    ch.y_axis.minorGridlines = ChartLines()
    ch.x_axis.minorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E2E2E2'))
    ch.y_axis.minorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E2E2E2'))
    ch.x_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='A6A6A6'))
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='A6A6A6'))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.number_format = '0'
    ch.y_axis.number_format = '0'


def pts_of(r):
    out = []
    for k in sorted(r['hmax'], key=lambda x: int(x)):
        W, v = int(k), r['hmax'][k]
        out.append((W, None if v is None else (H_TOP if v == 'なし' else v), v == 'なし'))
    return out


def graph_range(r):
    ok = [W for W, h, _ in pts_of(r) if h is not None]
    if not ok:
        return 0, 300
    lo = 0 if (r.get('wMin') or 0) < 200 else max(0, int((r['wMin'] - 30) // 50) * 50)
    hi = 300 if (r.get('wMin') or 0) < 200 else min(max(ok), 400)
    return lo, hi


wb = Workbook()
ov = wb.active
ov.title = '全部の型'
first_data_col = 30          # 表のデータ置き場（印刷範囲の外）

# ---------------------------------------------------------------- 1枚目：全部の型
title(ov, 'コの字曲げ 早見グラフ ― 型ごとの 低いほうの立上りの上限（全部の型）　株式会社高橋鉄骨',
      '線＝その底Wで普通に曲げられる低いほうの立上りの上限（外寸）。線にマウスを当てると数字が出ます。型ごとの詳しい数字は2枚目から（1型1枚）。'
      f'左右の立上りが違うときは低いほうで見ます（高いほうは計算では上限なし。ただし高すぎると機械から抜けないことがある）。計算は {TODAY} 時点・ヤゲン904061・中間板標準・普通に曲げたとき。',
      20)
Ws_all = list(range(5, 405, 5))
ov.cell(4, first_data_col, '底W')
for i, W in enumerate(Ws_all):
    ov.cell(5 + i, first_data_col, W)
ch = ScatterChart()
ch.style = 2
ch.width, ch.height = 40, 21
for j, r in enumerate(rows):
    c = first_data_col + 1 + j
    ov.cell(4, c, f"V{r['V']} t{r['t']}{'' if r['base'] else '（小さいV）'}")
    for i, W in enumerate(Ws_all):
        v = r['hmax'].get(str(W))
        ov.cell(5 + i, c, None if v is None else (H_TOP if v == 'なし' else v))
    xs = Reference(ov, min_col=first_data_col, min_row=5, max_row=4 + len(Ws_all))
    ys = Reference(ov, min_col=c, min_row=4, max_row=4 + len(Ws_all))
    s = Series(ys, xs, title_from_data=True)
    s.marker = Marker(symbol='none')
    s.graphicalProperties = line_props(COLORS[j % len(COLORS)], 1.75)
    s.smooth = False
    ch.series.append(s)
axes(ch, 0, 400, 300)
ch.legend.position = 'r'
ov.add_chart(ch, 'A3')
for rr in range(3, 50):
    ov.row_dimensions[rr].height = 15   # 行の高さを決めておく（既定の高さは Excel の設定で変わり、グラフが印刷で切れる）
a3(ov, 'A1:X47')

# ---------------------------------------------------------------- 2枚目から：1型1枚
for gi, r in enumerate(rows, start=1):
    tag = '基準金型' if r['base'] else '小さいV'
    ws = wb.create_sheet(f"t{r['t']} V{r['V']}")
    G_SPAN = 15
    t0 = G_SPAN + 2
    title(ws, f"コの字曲げ 早見グラフ {gi}／{len(rows)}　V{r['V']}　{r['sel'].split(':')[1]}・{r['machine']}　鉄 t{r['t']}（{tag}）",
          '線＝普通に曲げられる低いほうの立上りの上限（5mm ごとの点。点にマウスを当てると底W・立上りが出ます）。数字は10mm ごと、赤い点＝底Wを広げたのに上限が下がる所。'
          '右の表＝5mm ごとに、その底Wで曲げられる立上り（下限〜上限）。緑＝いちばん小さい断面。'
          f'左右の立上りが違うときは低いほうで見ます（高いほうは計算では上限なし。ただし高すぎると機械から抜けないことがある）。計算は {TODAY} 時点。',
          t0 + 5)
    pts = pts_of(r)
    lo, hi = graph_range(r)
    # グラフのデータ（印刷範囲の外の列に置く）
    dc = 40
    ws.cell(3, dc, '底W'); ws.cell(3, dc + 1, '上限'); ws.cell(3, dc + 2, '10mmごとの数字'); ws.cell(3, dc + 3, '下がる所')
    ws.cell(3, dc + 4, '最小フランジ'); ws.cell(3, dc + 5, 'いちばん小さい')
    shown = [(W, h, top) for W, h, top in pts if lo <= W <= hi]
    prev = None
    for i, (W, h, top) in enumerate(shown):
        rr = 4 + i
        ws.cell(rr, dc, W)
        ws.cell(rr, dc + 1, h)
        ws.cell(rr, dc + 2, h if (h is not None and W % 10 == 0 and not top) else None)
        ws.cell(rr, dc + 3, h if (h is not None and prev is not None and h < prev) else None)
        prev = h
    n = len(shown)
    ymax = min(300, max(150, (int(max([h for _, h, _ in shown if h is not None] + [100])) // 50 + 1) * 50))
    # 最小フランジの横線・いちばん小さい点
    ws.cell(4, dc + 6, lo); ws.cell(4, dc + 4, r['minOut'])
    ws.cell(5, dc + 6, hi); ws.cell(5, dc + 4, r['minOut'])
    if r.get('wMin'):
        ws.cell(4, dc + 7, r['wMin']); ws.cell(4, dc + 5, r['hLow'])

    ch = ScatterChart()
    ch.style = 2
    ch.width, ch.height = 26.5, 21.5
    X = Reference(ws, min_col=dc, min_row=4, max_row=3 + n)
    s1 = Series(Reference(ws, min_col=dc + 1, min_row=4, max_row=3 + n), X, title='上限（5mmごと）')
    s1.marker = Marker(symbol='circle', size=4)
    s1.marker.graphicalProperties = GraphicalProperties(solidFill='2A78D6')
    s1.marker.graphicalProperties.line.solidFill = '2A78D6'
    s1.graphicalProperties = line_props('2A78D6', 2.25)
    s1.smooth = False
    s2 = Series(Reference(ws, min_col=dc + 2, min_row=4, max_row=3 + n), X, title='10mmごとの数字')
    s2.marker = Marker(symbol='none')
    s2.graphicalProperties = line_props(None, none=True)
    s2.dLbls = DataLabelList()
    s2.dLbls.showVal = True
    s2.dLbls.showSerName = False
    s2.dLbls.showCatName = False
    s2.dLbls.showLegendKey = False
    s2.dLbls.position = 't'
    s3 = Series(Reference(ws, min_col=dc + 3, min_row=4, max_row=3 + n), X, title='下がる所')
    s3.marker = Marker(symbol='circle', size=8)
    s3.marker.graphicalProperties = GraphicalProperties(solidFill='C62F2F')
    s3.marker.graphicalProperties.line.solidFill = 'C62F2F'
    s3.graphicalProperties = line_props(None, none=True)
    s4 = Series(Reference(ws, min_col=dc + 4, min_row=4, max_row=5), Reference(ws, min_col=dc + 6, min_row=4, max_row=5), title=f"最小フランジ {r['minOut']}")
    s4.marker = Marker(symbol='none')
    s4.graphicalProperties = line_props('888888', 1.25, dash='dash')
    ch.series.extend([s1, s2, s3, s4])
    if r.get('wMin'):
        s5 = Series(Reference(ws, min_col=dc + 5, min_row=4, max_row=4), Reference(ws, min_col=dc + 7, min_row=4, max_row=4),
                    title=f"いちばん小さい 立上り{r['hLow']}×底W{r['wMin']}")
        s5.marker = Marker(symbol='square', size=10)
        s5.marker.graphicalProperties = GraphicalProperties(solidFill='0A7D0A')
        s5.marker.graphicalProperties.line.solidFill = '0A7D0A'
        s5.graphicalProperties = line_props(None, none=True)
        ch.series.append(s5)
    axes(ch, lo, hi, ymax)
    ch.legend.position = 't'
    ws.add_chart(ch, 'A3')

    # 右の表：底W（5mm＋いちばん狭い底W）と立上り（下限〜上限）
    for c in range(1, G_SPAN + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.column_dimensions[get_column_letter(G_SPAN + 1)].width = 2
    vals = [(int(k), r['hmax'][k]) for k in sorted(r['hmax'], key=lambda x: int(x)) if lo <= int(k) <= max(hi, lo)]
    vals = [(W, v) for W, v in vals if W <= hi or v not in (None,)]
    per = -(-len(vals) // 3)
    ws.merge_cells(start_row=3, start_column=t0, end_row=3, end_column=t0 + 5)
    sm = ws.cell(3, t0, f"いちばん小さい断面：立上り {r['hLow']} × 底W {r['wMin']}" if r.get('wMin') else 'この型は曲がらない')
    sm.font = Font(name=FONT, size=11.5, bold=True, color='0A7D0A')
    for b in range(3):
        ws.column_dimensions[get_column_letter(t0 + 2 * b)].width = 7.5
        ws.column_dimensions[get_column_letter(t0 + 2 * b + 1)].width = 10
        put(ws, 4, t0 + 2 * b, '底W', bold=True, size=10, fill=HEAD)
        put(ws, 4, t0 + 2 * b + 1, '立上り', bold=True, size=10, fill=HEAD)
    d = dict(vals)
    for idx, (W, v) in enumerate(vals):
        b, k = divmod(idx, per)
        rr = 5 + k
        first = W == r.get('wMin')
        pv = d.get(W - 5)
        dip = isinstance(v, (int, float)) and isinstance(pv, (int, float)) and v < pv
        put(ws, rr, t0 + 2 * b, W, size=10.5, bold=True, fill=MINF if first else WF)
        if v is None:
            put(ws, rr, t0 + 2 * b + 1, '×', size=10.5, color='999999')
        elif v == 'なし':
            put(ws, rr, t0 + 2 * b + 1, f"{r['hLow']}〜なし", size=10.5, color='0A7D0A')
        else:
            put(ws, rr, t0 + 2 * b + 1, (f"{v}だけ" if v == r['hLow'] else f"{r['hLow']}〜{v}"), size=11, bold=dip or first,
                color='C62F2F' if dip else ('0A7D0A' if first else None), fill=MINF if first else None)
    for rr in range(4, 4 + max(per, 44) + 1):
        ws.row_dimensions[rr].height = 15
    a3(ws, f'A1:{get_column_letter(t0 + 5)}{4 + max(per, 44)}')

wb.save(out_path)
print('saved', out_path)
