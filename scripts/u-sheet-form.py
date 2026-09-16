# -*- coding: utf-8 -*-
# コの字曲げ 実績記入シート（A3縦1枚）を作り、シミュレーション結果を入れる。
#   ・行（下型・材質・板厚・最小外寸）は Z曲げ実績記入シート と同じ並び。
#   ・うすい黄＝現場で実際の値を書き込む欄（空欄のまま）。
#   ・うすい緑＝シミュレーションの結果（scripts/u-sheet-fill.mjs）。
#
# 使い方: python scripts/u-sheet-form.py <Z曲げ実績記入シート.xlsx> <出力xlsx>
import sys, os, json, subprocess, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.drawing.image import Image as XLImage
from openpyxl.comments import Comment
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

zsrc, dst = sys.argv[1], sys.argv[2]
today = datetime.date.today().isoformat()
TMP = os.environ.get("TEMP", ".")

# ---------------------------------------------------------------- 行データ（Zシートから）
zs = openpyxl.load_workbook(zsrc, data_only=True).active
records = []
for r in range(8, 52):
    a = zs[f"A{r}"].value
    if a is None:
        continue
    if str(a).startswith("材質："):
        records.append(("sec", str(a))); continue
    V, t = zs[f"B{r}"].value, zs[f"E{r}"].value
    if not isinstance(V, (int, float)):
        continue
    records.append(("row", {"key": r, "die": a, "V": int(V), "R": zs[f"C{r}"].value, "mat": zs[f"D{r}"].value,
                            "t": float(t), "minOut": zs[f"F{r}"].value}))

rows_in = [{"row": x["key"], "V": x["V"], "mat": x["mat"], "t": x["t"], "minOut": x["minOut"]}
           for k, x in records if k == "row"]
p = os.path.join(TMP, "u_rows.json")
json.dump(rows_in, open(p, "w", encoding="utf-8"), ensure_ascii=False)
res = subprocess.run(["node", "scripts/u-sheet-fill.mjs", p], capture_output=True, text=True, encoding="utf-8")
if res.returncode != 0:
    print(res.stderr); sys.exit(1)
sim = {o["row"]: o for o in json.loads(res.stdout)}
json.dump(list(sim.values()), open(os.path.join(TMP, "u_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# ---------------------------------------------------------------- 寸法の取り方の図
for fp in ["C:/Windows/Fonts/YuGothM.ttc", "C:/Windows/Fonts/meiryo.ttc"]:
    if os.path.exists(fp):
        font_manager.fontManager.addfont(fp)
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=fp).get_name(); break
fig, ax = plt.subplots(figsize=(4.4, 3.2), dpi=200)
T = 6; H = 50; W = 90
outer = [(0, H), (0, 0), (W, 0), (W, H), (W - T, H), (W - T, T), (T, T), (T, H)]
ax.fill([q[0] for q in outer], [q[1] for q in outer], facecolor="#f6c9c9", edgecolor="#c0392b", lw=1.6)
def dim(x1, y1, x2, y2, txt, off, vertical=False):
    ax.annotate("", xy=(x1, y1), xytext=(x2, y2), arrowprops=dict(arrowstyle="<->", lw=0.9, color="#333"))
    if vertical:
        ax.text(x1 + off, (y1 + y2) / 2, txt, va="center", ha="left" if off > 0 else "right", fontsize=9)
    else:
        ax.text((x1 + x2) / 2, y1 + off, txt, ha="center", va="bottom" if off > 0 else "top", fontsize=9)
dim(0, -9, W, -9, "溝幅 W（外寸）", -2)
dim(-9, 0, -9, H, "フランジ H1\n（外寸）", -2, True)
dim(W + 9, 0, W + 9, H, "フランジ H2\n（外寸）", 2, True)
for x in (0, W):
    ax.plot([x, x], [-12, 0], color="#999", lw=0.6, ls="--")
for x in (-12, W + 12):
    pass
ax.plot([-12, 0], [0, 0], color="#999", lw=0.6, ls="--"); ax.plot([-12, 0], [H, H], color="#999", lw=0.6, ls="--")
ax.plot([W, W + 12], [0, 0], color="#999", lw=0.6, ls="--"); ax.plot([W, W + 12], [H, H], color="#999", lw=0.6, ls="--")
ax.text(W / 2, H + 14, "コの字曲げの寸法の取り方", ha="center", fontsize=11, fontweight="bold")
ax.text(W / 2, -26, "※ すべて板の外側で測る", ha="center", fontsize=8.5, color="#555")
ax.set_xlim(-44, W + 44); ax.set_ylim(-32, H + 22); ax.set_aspect("equal"); ax.axis("off")
diag = os.path.join(TMP, "u_diagram.png")
fig.savefig(diag, bbox_inches="tight", facecolor="white"); plt.close(fig)

# ---------------------------------------------------------------- シート
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "コの字曲げ実績"
F_KEY = PatternFill("solid", fgColor="EDF0F5"); F_HKEY = PatternFill("solid", fgColor="D6DCE5")
F_REF = PatternFill("solid", fgColor="DEEAF6"); F_HREF = PatternFill("solid", fgColor="BDD7EE")
F_IN = PatternFill("solid", fgColor="FFF2CC"); F_IN2 = PatternFill("solid", fgColor="FCE4A6"); F_HIN = PatternFill("solid", fgColor="F7DF8E")
F_SIM = PatternFill("solid", fgColor="E2EFDA"); F_HSIM = PatternFill("solid", fgColor="C6E0B4")
F_SEC = PatternFill("solid", fgColor="D9D9D9"); F_NOTE = PatternFill("solid", fgColor="F2F6FB")
thin = Side(style="thin", color="8EA9C1"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

COLS = [("A", "下型", 8), ("B", "V幅", 5.5), ("C", "曲げ\n内R", 6), ("D", "材質", 5.5), ("E", "板厚 t", 6.5),
        ("F", "最小フランジ\n外寸（表）", 11),
        ("G", "溝幅W\n最小", 9), ("H", "フランジH 上限\n（そのWのとき）", 13), ("I", "フランジH\n最大", 9), ("J", "備考", 18),
        ("K", "溝幅W 最小", 10), ("L", "当たる相手\n（W最小）", 14), ("M", "フランジH上限\nW=50", 11.5),
        ("N", "フランジH上限\nW=100", 11.5), ("O", "フランジH上限\nW=200", 11.5), ("P", "当たる相手\n（W=100のH上限）", 15),
        ("Q", "計算条件", 21)]
for c, _, w in COLS:
    ws.column_dimensions[c].width = w
LAST = "Q"

def put(coord, val, fill=None, font=None, align=CEN, border=True):
    cell = ws[coord]; cell.value = val
    if fill: cell.fill = fill
    if font: cell.font = font
    cell.alignment = align
    if border: cell.border = BOX
    return cell

ws.merge_cells(f"A1:{LAST}1")
put("A1", "コの字曲げ 実績記入シート　―　下型14型", font=Font(bold=True, size=20), align=Alignment(horizontal="left", vertical="center"), border=False)
ws.row_dimensions[1].height = 30
ws.merge_cells(f"A2:{LAST}2")
put("A2", "株式会社高橋鉄骨　曲げ可否判断シート用　／　うすい青＝参考（表の値）　うすい黄＝実際の値を書き込む欄　うすい緑＝シミュレーションの結果"
    f"（{today} 計算）", font=Font(size=10.5, color="404040"), align=Alignment(horizontal="left", vertical="center"), border=False)
ws.row_dimensions[2].height = 17

ws.merge_cells("A3:F10"); ws.merge_cells(f"G3:{LAST}10")
put("A3", None, fill=F_NOTE, border=False)
put("G3", "書き方　①寸法は左の絵のとおり。フランジH1・H2と溝幅Wはすべて外寸。　②うすい黄の欄に、実際に曲げた値を書く。分かるところだけでよい。\n"
          "　　　③上限が無いなら「なし」と書く。　④表に無い材質・板厚は「その他」の行へ。\n"
          "見方　・溝幅W最小 … 2か所目を曲げるとき、先に立てたフランジがヤゲン・中間板に当たらずに曲げられる最小の溝幅。\n"
          "　　　・フランジH上限 … その溝幅で、先に立てたフランジがヤゲン・中間板・機械上部に当たらない最大の高さ。\n"
          "　　　・緑の欄は計算の値です。実際の値と違ったら、黄の欄に実際の値を書いてください。次の計算に反映します。",
    fill=F_NOTE, font=Font(size=10), align=Alignment(horizontal="left", vertical="center", wrap_text=True), border=False)
for r in range(3, 11):
    ws.row_dimensions[r].height = 19
img = XLImage(diag); img.width, img.height = 250, 182; img.anchor = "A3"; ws.add_image(img)

HR1, HR2 = 12, 13
ws.merge_cells(f"A{HR1}:F{HR1}"); ws.merge_cells(f"G{HR1}:J{HR1}"); ws.merge_cells(f"K{HR1}:{LAST}{HR1}")
put(f"A{HR1}", "金型・材質・板厚", fill=F_HKEY, font=Font(bold=True, size=11))
put(f"G{HR1}", "記 入 ── 実際の値", fill=F_HIN, font=Font(bold=True, size=11))
put(f"K{HR1}", f"シミュレーション（{today}・ヤゲン904061・中間板標準・V.dxf で確認した土台）", fill=F_HSIM, font=Font(bold=True, size=11))
ws.row_dimensions[HR1].height = 20
for c, title, _ in COLS:
    fill = F_HKEY if c <= "E" else F_HREF if c == "F" else F_HIN if c <= "J" else F_HSIM
    put(f"{c}{HR2}", title, fill=fill, font=Font(bold=True, size=9.5))
ws.row_dimensions[HR2].height = 42

row = HR2 + 1
merge_start = merge_key = None
def flush(end):
    if merge_start and end >= merge_start + 1:
        for c in "ABC":
            ws.merge_cells(f"{c}{merge_start}:{c}{end}")

def val(v, why):
    return v if v is not None else (why or "—")

for kind, rec in records:
    if kind == "sec":
        flush(row - 1); merge_start = merge_key = None
        ws.merge_cells(f"A{row}:{LAST}{row}")
        put(f"A{row}", rec, fill=F_SEC, font=Font(bold=True, size=10.5), align=Alignment(horizontal="left", vertical="center"))
        ws.row_dimensions[row].height = 18; row += 1; continue
    key = (rec["die"], rec["mat"])
    if key != merge_key:
        flush(row - 1); merge_start, merge_key = row, key
    put(f"A{row}", rec["die"], fill=F_KEY, font=Font(bold=True, size=10.5))
    put(f"B{row}", rec["V"], fill=F_KEY, font=Font(size=10))
    put(f"C{row}", rec["R"], fill=F_KEY, font=Font(size=10))
    put(f"D{row}", rec["mat"], fill=F_KEY, font=Font(size=10))
    put(f"E{row}", rec["t"], fill=F_KEY, font=Font(size=10))
    put(f"F{row}", rec["minOut"], fill=F_REF, font=Font(size=10, color="404040"))
    for c in "GHI":
        put(f"{c}{row}", None, fill=F_IN2 if c == "G" else F_IN)
    put(f"J{row}", None, fill=F_IN, align=LEFT)
    o = sim.get(rec["key"], {})
    if "why" in o and "Wmin" not in o:
        for c in "KLMNOP":
            put(f"{c}{row}", "—", fill=F_SIM, font=Font(size=9.5, color="808080"))
        put(f"Q{row}", o["why"], fill=F_SIM, font=Font(size=8.5, color="404040"), align=LEFT)
    else:
        put(f"K{row}", val(o["Wmin"], o["WminWhy"]), fill=F_SIM, font=Font(size=10.5, bold=True, color="1F5F2F"))
        put(f"L{row}", o["WminHit"] or "—", fill=F_SIM, font=Font(size=8.5, color="404040"))
        for col, W in (("M", "50"), ("N", "100"), ("O", "200")):
            h = o["atW"][W]
            put(f"{col}{row}", val(h["v"], h["why"]), fill=F_SIM,
                font=Font(size=10.5, bold=h["v"] is not None, color="1F5F2F" if h["v"] is not None else "808080"))
        put(f"P{row}", o["atW"]["100"]["hit"] or "—", fill=F_SIM, font=Font(size=8.5, color="404040"))
        put(f"Q{row}", f"片伸び{o['nobi']}／W最小はH={o['H0']}", fill=F_SIM, font=Font(size=8.5, color="404040"), align=LEFT)
    ws.row_dimensions[row].height = 24
    row += 1
flush(row - 1)

ws.merge_cells(f"A{row}:{LAST}{row}")
put(f"A{row}", "その他（表に行が無い組合せ・特殊なケース）　下型・材質・板厚も書いてください", fill=F_SEC,
    font=Font(bold=True, size=10.5), align=Alignment(horizontal="left", vertical="center"))
ws.row_dimensions[row].height = 18; row += 1
for _ in range(4):
    for c, _, _ in COLS:
        fill = F_KEY if c <= "E" else F_REF if c == "F" else F_IN if c <= "J" else F_SIM
        put(f"{c}{row}", None, fill=fill)
    ws.row_dimensions[row].height = 24; row += 1

row += 1
notes = [
    "※ 緑の欄は bending-simulator.jsx の干渉判定そのもので計算しました（scripts/u-sheet-fill.mjs）。実測ではありません。",
    "※ 条件：上型ヤゲン904061・中間板標準、H1とH2は同じ高さ、2か所とも90°。突き当て（左右）は工程ごとに通るほうを使っています。",
    "※ 溝幅W最小は、フランジを通る最短（折り曲げ表の最小外寸＋5mm、V肩に届く長さ以上）で計算。高いフランジではW最小は大きくなります。",
    "※ フランジH上限（W=50・100・200）は、その溝幅で先に立てたフランジが当たらない最大の高さ。「W不足」はその幅が溝幅W最小より狭いという意味です。",
    "※ 「なし」は400mmまで伸ばしても当たらないという意味です。「当たる相手」は、それを超えたときに最初に当たる部品です。",
    "※ Z曲げの突き合わせでは V12・V16・V20・V25 のシミュが実績より甘く出ました。コの字でも同じ型は実績で必ず確かめてください。",
]
for n in notes:
    ws.merge_cells(f"A{row}:{LAST}{row}")
    put(f"A{row}", n, font=Font(size=9.5, color="404040"), align=Alignment(horizontal="left", vertical="center"), border=False)
    ws.row_dimensions[row].height = 15; row += 1

ws.page_setup.paperSize = 8
ws.page_setup.orientation = "portrait"
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.print_area = f"A1:{LAST}{row - 1}"
ws.print_title_rows = f"{HR1}:{HR2}"
ws.print_options.horizontalCentered = True
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.35
ws.freeze_panes = f"G{HR2 + 1}"
ws.sheet_view.showGridLines = False
wb.save(dst)
print("書き出し", dst, "／ 最終行", row - 1, "／ 列幅合計", sum(w for _, _, w in COLS))
