# コの字曲げ 実績記入シート2（中押し入り）を A3 縦1枚で作る。
#   計算値は scripts/u-sheet2-sim.mjs の出力（JSON）を使う。
#   使い方: python scripts/u-sheet2-form.py <sim.json> <出力.xlsx> <図.png>
import json
import sys
from datetime import date

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import rcParams
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

sim_path, out_path, fig_path = sys.argv[1], sys.argv[2], sys.argv[3]
rows = json.load(open(sim_path, encoding='utf-8'))

R_OF = {8: 1.3, 12: 2, 16: 2.6, 20: 3.3, 25: 4, 32: 5, 40: 6.5, 50: 8, 63: 10, 80: 13, 100: 16, 125: 20, 160: 26}
FONT = 'Meiryo UI'

# ---------------------------------------------------------------- 図（寸法の取り方と中押しの手順）
rcParams['font.family'] = ['Meiryo', 'Yu Gothic', 'MS Gothic']
fig, axs = plt.subplots(1, 4, figsize=(11.5, 2.6), gridspec_kw={'width_ratios': [1.25, 1, 1, 1]})
for ax in axs:
    ax.set_aspect('equal'); ax.axis('off')
lw = 5
def arrow(ax, x1, y1, x2, y2, txt, tx, ty, **kw):
    ax.annotate('', xy=(x1, y1), xytext=(x2, y2), arrowprops=dict(arrowstyle='<->', color='#555', lw=1.2))
    ax.text(tx, ty, txt, ha='center', va='center', fontsize=12, fontweight='bold', **kw)
a = axs[0]
a.plot([0, 0, 100, 100], [80, 0, 0, 80], color='#111', lw=lw, solid_joinstyle='miter')
arrow(a, -14, 0, -14, 80, '', 0, 0); a.text(-30, 40, 'H1', fontsize=13, fontweight='bold', ha='center', va='center')
arrow(a, 114, 0, 114, 80, '', 0, 0); a.text(130, 40, 'H2', fontsize=13, fontweight='bold', ha='center', va='center')
arrow(a, 0, -14, 100, -14, '', 0, 0); a.text(50, -27, '底W', fontsize=13, fontweight='bold', ha='center', va='center')
a.text(50, 45, 'すべて外寸', fontsize=10, ha='center', color='#444')
a.set_xlim(-45, 145); a.set_ylim(-40, 95)
a.set_title('寸法の取り方', fontsize=11)
# 中押しの手順
import math
def step(ax, peak, flange_deg, title, note):
    # 底は左右の半分が peak 度ずつ上がる（への字）。立上りは各半分に対して90°
    half = 50
    th = math.radians(peak)
    cx, cy = 0, 0
    L = (cx - half * math.cos(th), cy - half * math.sin(th))
    Rr = (cx + half * math.cos(th), cy - half * math.sin(th))
    fl = 55
    phi = math.radians(flange_deg)
    Lt = (L[0] - fl * math.sin(th) * (1 if flange_deg else 0) + (0 if flange_deg else -fl), L[1] + fl * math.cos(th) * (1 if flange_deg else 0))
    Rt = (Rr[0] + fl * math.sin(th) * (1 if flange_deg else 0) + (0 if flange_deg else fl), Rr[1] + fl * math.cos(th) * (1 if flange_deg else 0))
    xs = [Lt[0], L[0], cx, Rr[0], Rt[0]]; ys = [Lt[1], L[1], cy, Rr[1], Rt[1]]
    ax.plot(xs, ys, color='#111', lw=lw, solid_joinstyle='miter')
    ax.set_xlim(-95, 95); ax.set_ylim(-45, 80)
    ax.set_title(title, fontsize=11)
    ax.text(0, -38, note, ha='center', fontsize=9, color='#444')
b = axs[1]; step(b, 12, 0, '① 底の真ん中を への字に', '真ん中を少し曲げる（への字の角度）')
b.annotate('', xy=(0, 2), xytext=(0, 28), arrowprops=dict(arrowstyle='->', color='#2a78d6', lw=2))
c = axs[2]; step(c, 12, 90, '② 両サイドを 90°', '立上りはまだ少し外へ開いている')
d = axs[3]; step(d, 0, 90, '③ 真ん中を押して戻す（中押し）', 'ヤゲンがコの字の内側に入る')
d.plot([-5, -5, 5, 5], [70, 4, 4, 70], color='#2a78d6', lw=2)
d.annotate('', xy=(0, 1), xytext=(0, -10), arrowprops=dict(arrowstyle='-', color='white'))
plt.tight_layout()
plt.savefig(fig_path, dpi=160)
plt.close()

# ---------------------------------------------------------------- シート
wb = Workbook()
ws = wb.active
ws.title = 'コの字曲げ実績'

thin = Side(style='thin', color='999999')
med = Side(style='medium', color='555555')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def fill(hex6): return PatternFill('solid', fgColor=hex6)
C_HEAD, C_KEY, C_REF, C_CALC, C_IN, C_NAKA_H, C_SEC = 'BFBFBF', 'EAEAEA', 'DCE6F1', 'E2EFDA', 'FFF2CC', 'FCE4D6', 'D9D9D9'
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)

widths = {'A': 7, 'B': 5, 'C': 5.5, 'D': 5, 'E': 5.5, 'F': 8,
          'G': 8, 'H': 9, 'I': 8, 'J': 9, 'K': 8, 'L': 9, 'M': 8, 'N': 9,
          'O': 8, 'P': 9, 'Q': 8, 'R': 9, 'S': 9, 'T': 9.5, 'U': 24}
for k, v in widths.items():
    ws.column_dimensions[k].width = v
LAST = 'U'

def put(cell, v, bg=None, bold=False, size=10, al=center, color='000000', italic=False):
    c = ws[cell]
    c.value = v
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    c.alignment = al
    if bg: c.fill = fill(bg)
    return c
def box(rng, bg=None):
    if ':' not in rng:
        rng = f'{rng}:{rng}'
    for row in ws[rng]:
        for c in row:
            c.border = BOX
            if bg: c.fill = fill(bg)
def merge(rng, v, **kw):
    ws.merge_cells(rng)
    put(rng.split(':')[0], v, **kw)
    box(rng, kw.get('bg'))

# 見出し
ws.row_dimensions[1].height = 28
ws.merge_cells('A1:U1')
put('A1', 'コの字曲げ 実績記入シート（中押し入り）　―　下型14型', size=16, bold=True, al=left)
ws.row_dimensions[2].height = 18
put('A2', f'株式会社高橋鉄骨　曲げ可否判断シート用　／　作成 {date.today():%Y-%m-%d}　／　'
          '色：灰＝型と板（さわらない）　青＝折り曲げ表　緑＝計算の値（さわらない）　黄＝実際の値を書く欄', size=10, al=left)
ws.merge_cells('A2:U2')

# 図と書き方
for r in range(3, 10):
    ws.row_dimensions[r].height = 22.5
img = XLImage(fig_path)
img.width, img.height = 880, 199
ws.add_image(img, 'A3')
merge('Q3:U9',
      '書き方\n'
      '① 寸法は左の絵のとおり、すべて外寸。\n'
      '② 緑＝計算の値。その右の黄の欄に、実際に曲げた値を書く。分かる所だけでよい。\n'
      '③ 計算どおりなら「〃」、上限が無いなら「なし」。\n'
      '④ 中押し＝③の手順（への字 → 両サイド90° → 真ん中を押して戻す）。\n'
      '⑤ 表に無い板厚は「その他」、品物ごとの記録は下の「実例」へ。',
      bg='F2F6FB', size=9.5, al=Alignment(horizontal='left', vertical='top', wrap_text=True))

# 表の見出し（3段）
H1, H2, H3 = 11, 12, 13
ws.row_dimensions[H1].height = 20
ws.row_dimensions[H2].height = 40
ws.row_dimensions[H3].height = 16
merge(f'A{H1}:F{H1}', '型・材質・板厚', bg=C_HEAD, bold=True)
merge(f'G{H1}:N{H1}', '普通に曲げる（2か所とも90°）', bg=C_HEAD, bold=True)
merge(f'O{H1}:T{H1}', '中押し（への字 → 両サイド90° → 中押し）', bg=C_NAKA_H, bold=True)
merge(f'U{H1}:U{H3}', '備考\n（当たった所・\nやり方など）', bg=C_IN, bold=True, size=9)
for col, name in zip('ABCDE', ['下型', 'V幅', '曲げ\n内R', '材質', '板厚\nt']):
    merge(f'{col}{H2}:{col}{H3}', name, bg=C_KEY, bold=True, size=9)
merge(f'F{H2}:F{H3}', '最小\nフランジ\n外寸(表)', bg=C_REF, bold=True, size=8.5)
pairs = [('G', '底W 最小\n（立上り最短のとき）'), ('I', '立上りH 上限\n底W＝50'), ('K', '立上りH 上限\n底W＝100'),
         ('M', '立上りH 上限\n底W＝200'), ('O', '底W 最小\n（立上りH＝100）'), ('Q', '立上りH 上限\n（底W＝100）')]
for col, name in pairs:
    col2 = chr(ord(col) + 1)
    bg = C_NAKA_H if col in 'OQ' else C_HEAD
    merge(f'{col}{H2}:{col2}{H2}', name, bg=bg, bold=True, size=9)
    put(f'{col}{H3}', '計算', bg=C_CALC, bold=True, size=8.5); box(f'{col}{H3}')
    put(f'{col2}{H3}', '実際', bg=C_IN, bold=True, size=8.5); box(f'{col2}{H3}')
merge(f'S{H2}:S{H3}', 'への字の\n角度(°)', bg=C_IN, bold=True, size=8.5)
merge(f'T{H2}:T{H3}', '戻した後の\n底の反り\n(mm)', bg=C_IN, bold=True, size=8.5)

# データ
r = H3 + 1
first_data = r
def section(label):
    global r
    ws.row_dimensions[r].height = 15
    merge(f'A{r}:U{r}', label, bg=C_SEC, bold=True, size=9.5, al=left)
    r += 1
def num(v):
    return v
calc_cells = []
for mat in ['鉄', '縞']:
    section(f'材質：{mat}')
    for d in [x for x in rows if x['mat'] == mat]:
        ws.row_dimensions[r].height = 21
        V, t = d['V'], d['t']
        for col, v, bg in [('A', f'V{V}', C_KEY), ('B', V, C_KEY), ('C', R_OF.get(V), C_KEY), ('D', mat, C_KEY), ('E', t, C_KEY),
                           ('F', d['minOut'], C_REF)]:
            put(f'{col}{r}', v, bg=bg); box(f'{col}{r}')
        # 中押しは底の真ん中をV溝で曲げるので、底の半分ずつが最小フランジ以上いる
        naka_w = max(d['nakaW100'], int(-(-2 * d['minOut'] // 1)))
        naka_h = d['nakaH100'] if 100 >= 2 * d['minOut'] else '不可'
        fx = lambda v: '不可' if v == 'W不足' else v
        vals = {'G': d['wMin'], 'I': fx(d['h50']), 'K': fx(d['h100']), 'M': fx(d['h200']), 'O': naka_w, 'Q': naka_h}
        for col, v in vals.items():
            c = put(f'{col}{r}', v, bg=C_CALC, size=10); box(f'{col}{r}')
            calc_cells.append(f'{col}{r}')
            col2 = chr(ord(col) + 1)
            put(f'{col2}{r}', None, bg=C_IN); box(f'{col2}{r}')
        for col in 'STU':
            put(f'{col}{r}', None, bg=C_IN, al=left if col == 'U' else center, size=9); box(f'{col}{r}')
        if naka_w < d['nakaRuleW']:
            ws[f'O{r}'].comment = Comment(f'計算の値。現場の決まり（内-内120mm）なら底W {d["nakaRuleW"]}mm 以上。'
                                          f'{naka_w}〜{d["nakaRuleW"] - 1}mm はまだ確かめていない。', '曲げシミュレーター')
        r += 1
section('その他（表に無い材質・板厚）')
for _ in range(3):
    ws.row_dimensions[r].height = 21
    for col in 'ABCDEF': put(f'{col}{r}', None, bg=C_KEY); box(f'{col}{r}')
    for col in 'GIKMOQ':
        put(f'{col}{r}', None, bg=C_CALC); box(f'{col}{r}')
        put(f'{chr(ord(col) + 1)}{r}', None, bg=C_IN); box(f'{chr(ord(col) + 1)}{r}')
    for col in 'STU': put(f'{col}{r}', None, bg=C_IN); box(f'{col}{r}')
    r += 1
last_data = r - 1

# 実際の値が計算より 3mm 以上ちがうとき、黄の欄をオレンジにする（Excelで入力したとき）
orange = PatternFill('solid', fgColor='F8CBAD')
for col in 'GIKMOQ':
    col2 = chr(ord(col) + 1)
    rng = f'{col2}{first_data}:{col2}{last_data}'
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(ISNUMBER({col2}{first_data}),ISNUMBER({col}{first_data}),ABS({col2}{first_data}-{col}{first_data})>=3)'], fill=orange))

# 品物ごとの実例
r += 1
ws.row_dimensions[r].height = 18
merge(f'A{r}:U{r}', 'コの字・中押しの実例（品物ごと。曲がった・曲がらなかった、どちらも書く）', bg=C_HEAD, bold=True, al=left)
r += 1
ws.row_dimensions[r].height = 30
heads = [('A', 'B', '日付'), ('C', 'D', '材質'), ('E', 'E', '板厚'), ('F', 'F', '下型'), ('G', 'H', '底W'), ('I', 'J', 'H1'),
         ('K', 'L', 'H2'), ('M', 'N', '曲げ方\n普通／中押し'), ('O', 'P', '結果\n○／✕'), ('Q', 'R', 'への字\n角度(°)'),
         ('S', 'T', '底の反り\n(mm)'), ('U', 'U', 'メモ（当たった所など）')]
for a_, b_, name in heads:
    merge(f'{a_}{r}:{b_}{r}', name, bg=C_KEY, bold=True, size=9)
r += 1
ex_first = r
for i in range(7):
    ws.row_dimensions[r].height = 21
    for a_, b_, _ in heads:
        if a_ != b_: ws.merge_cells(f'{a_}{r}:{b_}{r}')
        box(f'{a_}{r}:{b_}{r}', C_IN)
        ws[f'{a_}{r}'].alignment = center
    r += 1
# 記入例（1行目。灰色の斜体で、消して使う）
ex = ['2026/9/22', '鉄', 4.5, 'V25', 100, 100, 100, '中押し', '○', 160, 0.5, '記入例（消して使う）']
for (a_, _, _), v in zip(heads, ex):
    put(f'{a_}{ex_first}', v, bg=C_IN, italic=True, color='808080', size=9)

# 注記
r += 1
notes = [
    '※ 緑の計算の値は、曲げシミュレーター（ヤゲン904061・中間板標準・V.dxf で確認した台）の干渉判定です。実測ではありません。',
    '※ 普通に曲げる：立上りH1とH2は同じ高さで計算。「不可」＝その底Wでは最短の立上りでも曲げられない（底が狭い）。「なし」＝400mmまで当たらない。',
    '※ 中押し：押し切った瞬間（底は平ら・立上りは垂直）に、ヤゲン・中間板・ホルダが立上りの内側に入るかの計算。底の半分ずつが最小フランジ以上いる。',
    '※ 中押しの現場の決まりは「底の内-内120mm以上」。緑の値がそれより狭い所（セルに三角の印）は、まだ確かめていない値です。',
    '※ Excelで入力したとき、実際の値が計算と3mm以上ちがうと、その欄がオレンジになります。',
]
for n in notes:
    ws.row_dimensions[r].height = 15
    put(f'A{r}', n, size=9, al=left)
    ws.merge_cells(f'A{r}:U{r}')
    r += 1

# 印刷：A3 縦 1枚
ws.print_area = f'A1:U{r - 1}'
ws.print_title_rows = f'{H1}:{H3}'
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.orientation = 'portrait'
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.35, bottom=0.35, header=0.15, footer=0.15)
ws.print_options.horizontalCentered = True
ws.freeze_panes = f'G{first_data}'
ws.sheet_view.zoomScale = 90
wb.save(out_path)
print('saved', out_path, 'rows', r - 1)
