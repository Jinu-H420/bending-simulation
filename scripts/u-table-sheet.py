# コの字曲げ 早見表（A3縦）を作る。グラフ版の実績入力シートを、数字で細かく引けるようにした別ファイル。
#   縦＝底W（5mm刻み）、横＝型（1枚に8型）、ます目＝低いほうの立上りの上限（1mm）。× は曲がらない、「なし」は上限なし。
#   左右の立上りが違うときは低いほうで引く（高いほうは計算ではどの型も上限なし。u-table-sim.mjs の説明）。
#   使い方: node scripts/u-table-sim.mjs > ut.json
#           python scripts/u-table-sheet.py ut.json <出力.xlsx>
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

sim_path, out_path = sys.argv[1], sys.argv[2]
rows = json.load(open(sim_path, encoding='utf-8'))
rows.sort(key=lambda r: (r['t'], r['base']))   # 板厚の小さい順。同じ板厚なら V8 を基準金型の前に
FONT = 'Meiryo UI'
TODAY = date.today().isoformat()
PER_PAGE = 8

thin = Side(style='thin', color='999999')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='E8EBE6')
OKF = PatternFill('solid', fgColor='FFFFFF')
FREE = PatternFill('solid', fgColor='EAF6EA')     # 上限なし
NGF = PatternFill('solid', fgColor='EFEFEF')      # 曲がらない
WF = PatternFill('solid', fgColor='F4F1E8')       # 底W の列


def put(ws, r, c, v, *, bold=False, size=11, fill=None, color=None, align='center'):
    cell = ws.cell(r, c, v)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    cell.border = box
    if fill:
        cell.fill = fill
    return cell


wb = Workbook()
pages = [rows[i:i + PER_PAGE] for i in range(0, len(rows), PER_PAGE)]
for pi, page in enumerate(pages, start=1):
    ws = wb.active if pi == 1 else wb.create_sheet()
    ws.title = f'早見表{pi}'
    ncol = 1 + len(page)
    # 見出し
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    t = ws.cell(1, 1, f'コの字曲げ 早見表 {pi}／{len(pages)}　― 低いほうの立上りの上限（mm）　株式会社高橋鉄骨')
    t.font = Font(name=FONT, size=15, bold=True)
    ws.row_dimensions[1].height = 22
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    n = ws.cell(2, 1,
                '見方　左の底W の行と、型の列が交わる数字が、普通に曲げられる「低いほうの立上り」の上限（外寸）。× は曲がらない、「なし」は上限なし。'
                '左右の立上りが違うときは、低いほうで引いてください（後から曲げる高いほうは、計算ではどの型も上型に当たらない）。'
                'ただし高いほうが高すぎると、曲げた後に機械から抜けないことがあります（計算に入っていない）。'
                f'中押し・くの字は入れていません。計算は {TODAY} 時点・ヤゲン904061・中間板標準。V12〜V25 はZ曲げで実際より甘く出た型。')
    n.font = Font(name=FONT, size=9.5)
    n.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 52
    # 型の見出し（3行）
    put(ws, 3, 1, '型', bold=True, fill=HEAD)
    put(ws, 4, 1, 'いちばん小さい\n立上り×底W', bold=True, size=9, fill=HEAD)
    put(ws, 5, 1, '底W ↓', bold=True, fill=HEAD)
    for j, r in enumerate(page, start=2):
        tag = '基準' if r['base'] else '小さいV'
        put(ws, 3, j, f"V{r['V']}　t{r['t']}\n{r['sel'].split(':')[1]}・{r['machine']}\n（{tag}）", bold=True, size=10, fill=HEAD)
        put(ws, 4, j, f"{r['hLow']}×{r['wMin']}" if r.get('wMin') else '—', size=10.5, fill=HEAD)
        put(ws, 5, j, f"最小フランジ {r['minOut']}", size=9, fill=HEAD)
    ws.row_dimensions[3].height = 60   # 型の名前（型番・機械が折り返して4行になる）
    ws.row_dimensions[4].height = 26
    ws.row_dimensions[5].height = 16
    # 底W の範囲：このページの型で、どれかが曲がる所から、全部「なし」になる所まで
    Ws = sorted({int(w) for r in page for w in r['hmax'].keys()})
    used = [w for w in Ws if any(r['hmax'].get(str(w)) is not None for r in page)]
    w_lo, w_hi = (min(used), max(used)) if used else (Ws[0], Ws[-1])
    while w_hi > w_lo and all(r['hmax'].get(str(w_hi)) in ('なし', None) for r in page) and \
            all(r['hmax'].get(str(w_hi - 5)) in ('なし', None) for r in page):
        w_hi -= 5
    rr = 6
    for w in [x for x in Ws if w_lo - 5 <= x <= w_hi]:
        put(ws, rr, 1, w, bold=True, size=11, fill=WF)
        for j, r in enumerate(page, start=2):
            v = r['hmax'].get(str(w))
            if v is None:
                put(ws, rr, j, '×', size=10, fill=NGF, color='999999')
            elif v == 'なし':
                put(ws, rr, j, 'なし', size=10, fill=FREE, color='0a7d0a')
            else:
                # 前の行より下がったら（くぼみ）赤字にする
                prev = r['hmax'].get(str(w - 5))
                dip = isinstance(prev, (int, float)) and v < prev
                put(ws, rr, j, v, size=11, fill=OKF, color='c62f2f' if dip else None, bold=dip)
        ws.row_dimensions[rr].height = 15
        rr += 1
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=ncol)
    foot = ws.cell(rr, 1, '赤字＝底Wを広げたのに上限が下がる所（2か所目を曲げるとき、先に立てた立上りが上型の別の所に当たるため）。')
    foot.font = Font(name=FONT, size=9.5, color='c62f2f')
    ws.column_dimensions['A'].width = 12
    for j in range(2, ncol + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15
    # A3 縦 1枚
    ws.page_setup.paperSize = 8
    ws.page_setup.orientation = 'portrait'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.35, bottom=0.3, header=0.2, footer=0.2)
    ws.freeze_panes = ws.cell(6, 2)

wb.save(out_path)
print('saved', out_path)
