# -*- coding: utf-8 -*-
# Z曲げ実績記入シートの「実測以外の空欄」をシミュレーション値で埋める。
#   ・実測（I〜N に入っている値）には一切触らない。
#   ・空いている I〜M にシミュ値を入れる。青い斜体＋セルのコメント「シミュレーション値」で区別。
#   ・R〜X（シミュ列）は、いまの土台（V.dxf で確認済み）で計算し直して上書きする。
#   ・元のファイルは上書きせず、別名で保存する。
#
# 使い方: python scripts/z-sheet-fill.py <入力xlsx> <出力xlsx>
import sys, json, subprocess, os, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
ws = wb.active

rows = []
for r in range(9, 52):
    V, t, mat = ws[f"B{r}"].value, ws[f"E{r}"].value, ws[f"D{r}"].value
    if not isinstance(V, (int, float)) or not isinstance(t, (int, float)):
        continue
    num = lambda c: ws[f"{c}{r}"].value if isinstance(ws[f"{c}{r}"].value, (int, float)) else None
    rows.append({"row": r, "V": int(V), "t": float(t), "mat": mat, "minOut": num("F"),
                 "actI": num("I"), "actK": num("K")})

tmp = os.path.join(os.environ.get("TEMP", "."), "zfill_rows.json")
json.dump(rows, open(tmp, "w", encoding="utf-8"), ensure_ascii=False)
res = subprocess.run(["node", "scripts/z-sheet-fill.mjs", tmp], capture_output=True, text=True, encoding="utf-8")
if res.returncode != 0:
    print(res.stderr); sys.exit(1)
sims = {o["row"]: o for o in json.loads(res.stdout)}

SIM_FONT = Font(color="1F4E9E", italic=True)
SIM_FILL = PatternFill("solid", fgColor="E8F0FC")
today = datetime.date.today().isoformat()

def put_sim(cell, value, note):
    cell.value = value
    cell.font = SIM_FONT
    cell.fill = SIM_FILL
    cell.comment = Comment(f"シミュレーション値（{today}）\n{note}", "sim")

filled = 0
for r, o in sims.items():
    if "why" in o and "simI" not in o:
        reason = o.get("simKwhy") or o["why"]
        if ws[f"K{r}"].value in (None, ""):
            put_sim(ws[f"K{r}"], "算出できず", f"{reason}（両端フランジ {o.get('flangeOuter')} mm）")
            filled += 1
        ws[f"R{r}"] = "算出できず"; ws[f"U{r}"] = reason
        continue
    cond = f"{o['die']}／ヤゲン904061／{'段差S=実績' if o['Ssrc']=='実績' else '段差S=sim最小'} {o['S']}／相手フランジ{o['other']}"
    # --- 記入欄（実測が空のところだけ）
    targets = [
        ("K", o["simK"], o["simKwhy"] or "両端フランジ %s mm で段差を詰めた最小" % o["flangeOuter"]),
        ("I", o["simI"], o["simIwhy"] or "その段差SでのフランジA上限。" + cond),
        ("J", o["simJ"], o["simJwhy"] or "段差を300mmとったときのフランジA上限（なし＝400mmでも当たらない）"),
        ("L", o["simL"], o["simLwhy"] or "その段差SでのフランジB最小。" + cond),
        ("M", o["simM"], o["simMwhy"] or "その段差SでのフランジB上限（なし＝400mmでも当たらない）。" + cond),
    ]
    for col, val, note in targets:
        c = ws[f"{col}{r}"]
        if c.value not in (None, ""):
            continue                       # 実測は触らない
        put_sim(c, val if val is not None else note[:20], note)
        filled += 1
    # --- シミュ列（R〜X）を今の土台で計算し直す
    actK = next((x["actK"] for x in rows if x["row"] == r), None)
    actI = next((x["actI"] for x in rows if x["row"] == r), None)
    ws[f"R{r}"] = o["simK"] if o["simK"] is not None else "算出できず"
    ws[f"S{r}"] = max(v for v in (o["simK"], actK) if v is not None) if (o["simK"] is not None or actK is not None) else None
    ws[f"T{r}"] = round(actK - o["simK"], 1) if (actK is not None and o["simK"] is not None) else None
    ws[f"U{r}"] = cond
    ws[f"V{r}"] = o["simI"]
    if actI is not None and isinstance(o["simI"], (int, float)):
        d = round(actI - o["simI"], 1)
        ws[f"W{r}"] = d
        ws[f"X{r}"] = "ほぼ一致" if abs(d) <= 5 else ("シミュレーションが甘い（危険）" if d < 0 else "シミュレーションが厳しめ")
    else:
        ws[f"W{r}"] = None; ws[f"X{r}"] = None

ws["R6"] = f"シミュレーション（{today} 再計算・V.dxf で確認した土台）"
ws["V6"] = f"フランジA上限の突き合わせ（{today}）"

# 注記
last = ws.max_row + 2
notes = [
    f"※ {today}：実測の入っていない I〜M 列を、シミュレーション値で埋めました（青い斜体・セルにコメントあり）。実測値ではありません。",
    "※ 段差S最小（K）が空の行は、その sim 値を使って I・L・M を計算しています。実測の K がある行は実測の段差で計算しています。",
    "※ フランジA最大（J）は段差を300mmとったときの上限、「なし」は400mmまで当たらないという意味です。",
    "※ R〜X 列は、V.dxf で確認した土台（HG2203 は V8〜V25 共通の台）で計算し直しました。",
]
for i, n in enumerate(notes):
    ws[f"A{last + i}"] = n
    ws[f"A{last + i}"].font = Font(color="1F4E9E")

wb.save(dst)
print("書き出し", dst, "／ 記入したセル", filled)
json.dump(list(sims.values()), open(os.path.join(os.environ.get("TEMP", "."), "zfill_result.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
