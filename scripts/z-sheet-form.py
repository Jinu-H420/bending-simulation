# -*- coding: utf-8 -*-
# Z曲げ実績記入シートを、配布用の A3縦1枚の記入フォームとして作り直す。
# 元データは docs/Z曲げ_実績記入シート.xlsx（訂正済みの実績が入っているほう）。
# レイアウトは Z曲げ_実績記入シート_A3.pdf に合わせている。
import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.drawing.image import Image as XLImage

SRC = "docs/Z曲げ_実績記入シート.xlsx"
DST = "docs/Z曲げ_実績記入シート_A3.xlsx"
DIAGRAM = "docs/Z寸法の取り方.jpg"
DRAWN = {12, 25, 32, 40, 63}            # 断面の図面がある型

src = openpyxl.load_workbook(SRC, data_only=True).active


def g(c, r):
    v = src["%s%d" % (c, r)].value
    return "" if v is None else v


# 元シートの 8〜51 行から、行と区切りを読む
records = []
for r in range(8, 52):
    a = g("A", r)
    if a == "":
        continue
    if str(a).startswith("材質："):
        records.append(("sec", str(a)))
        continue
    records.append(("row", {k: g(k, r) for k in "ABCDEFGHIKN"}))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Z曲げ実績"

NAVY = "1F4E79"
F_REF = PatternFill("solid", fgColor="DEEAF6")    # 参考
F_IN = PatternFill("solid", fgColor="FFF2CC")     # 記入
F_IN2 = PatternFill("solid", fgColor="FCE4A6")    # 記入（段差S）
F_KEY = PatternFill("solid", fgColor="EDF0F5")    # 金型
F_HREF = PatternFill("solid", fgColor="BDD7EE")
F_HIN = PatternFill("solid", fgColor="F7DF8E")
F_HKEY = PatternFill("solid", fgColor="D6DCE5")
F_SEC = PatternFill("solid", fgColor="D9D9D9")
F_NOTE = PatternFill("solid", fgColor="F2F6FB")
thin = Side(style="thin", color="8EA9C1")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

COLS = [("A", "下型", 9), ("B", "V幅", 6), ("C", "曲げ\n内R", 6.5), ("D", "材質", 6),
        ("E", "板厚 t", 7), ("F", "最小フランジ\n外寸（表）", 13), ("G", "段差S 最小\n暫定", 12),
        ("H", "段差S 最小\n幾何", 12), ("I", "フランジA 上限\n（その段差Sのとき）", 17.5),
        ("J", "フランジA\n最大", 11), ("K", "段差S\n最小", 10), ("L", "フランジB\n最小", 11),
        ("M", "フランジB\n最大", 11), ("N", "備考", 26)]
for c, _, w in COLS:
    ws.column_dimensions[c].width = w
LAST = "N"


def put(coord, val, fill=None, font=None, align=CEN, border=True):
    c = ws[coord]
    c.value = val
    if fill:
        c.fill = fill
    if font:
        c.font = font
    c.alignment = align
    if border:
        c.border = BOX
    return c


# ── 見出し ─────────────────────────────────────────────
ws.merge_cells("A1:%s1" % LAST)
put("A1", "Z曲げ（段違い）実績記入シート　―　下型14型", font=Font(bold=True, size=20),
    align=Alignment(horizontal="left", vertical="center"), border=False)
ws.row_dimensions[1].height = 30
ws.merge_cells("A2:%s2" % LAST)
put("A2", "株式会社高橋鉄骨　曲げ可否判断シート用　／　うすい青＝参考（いまの計算値・触らない）　"
          "うすい黄＝実際の値を書き込む欄　　●＝断面の図面がある型",
    font=Font(size=10.5, color="404040"),
    align=Alignment(horizontal="left", vertical="center"), border=False)
ws.row_dimensions[2].height = 17

# 図（左）と書き方（右）
ws.merge_cells("A3:E10")
ws.merge_cells("F3:%s10" % LAST)
put("A3", None, fill=F_NOTE, border=False)
put("F3", "書き方　①寸法は左の絵のとおり。フランジA・Bは外寸、段差Sは外-外。　"
          "②左のうすい青「参考」を見ながら、右のうすい黄に実際の値を書く。分かるところだけでよい。\n"
          "　　　③上限が無いなら「なし」と書く。　④表に無い材質・板厚は「その他」の行へ。\n"
          "参考の中身　・最小フランジ 外寸（表）… 折り曲げ表の最小外寸。これより短いとV溝に落ちる。\n"
          "　　　・段差S 暫定 … 3×(曲げ内R＋板厚)。実績2件に合わせた式。14型ぜんぶ出せます。\n"
          "　　　・段差S 幾何 … 当たり判定から出した値。●の5型だけ。角をピン角で見るので小さめに出ます。\n"
          "　　　実際の最小は暫定と幾何の大きいほうになるはず。そこを確かめたいです。",
    fill=F_NOTE, font=Font(size=10),
    align=Alignment(horizontal="left", vertical="center", wrap_text=True), border=False)
for r in range(3, 11):
    ws.row_dimensions[r].height = 19
if os.path.exists(DIAGRAM):
    img = XLImage(DIAGRAM)
    img.width, img.height = 243, 150
    img.anchor = "A3"
    ws.add_image(img)

# ── 表の見出し2段 ───────────────────────────────────────
HR1, HR2 = 12, 13
ws.merge_cells("A%d:E%d" % (HR1, HR1))
ws.merge_cells("F%d:H%d" % (HR1, HR1))
ws.merge_cells("I%d:%s%d" % (HR1, LAST, HR1))
put("A%d" % HR1, "金型・材質・板厚", fill=F_HKEY, font=Font(bold=True, size=11))
put("F%d" % HR1, "参 考 ── いまの計算値", fill=F_HREF, font=Font(bold=True, size=11))
put("I%d" % HR1, "記 入 ── 実際の値", fill=F_HIN, font=Font(bold=True, size=11))
ws.row_dimensions[HR1].height = 20
for c, title, _ in COLS:
    fill = F_HKEY if c <= "E" else (F_HREF if c <= "H" else F_HIN)
    put("%s%d" % (c, HR2), title, fill=fill, font=Font(bold=True, size=10))
ws.row_dimensions[HR2].height = 40

# ── 本体 ───────────────────────────────────────────────
row = HR2 + 1
merge_start = None
merge_key = None


def flush(end):
    if merge_start and end >= merge_start + 1:
        for c in "ABC":
            ws.merge_cells("%s%d:%s%d" % (c, merge_start, c, end))


for kind, rec in records:
    if kind == "sec":
        flush(row - 1)
        merge_start = None
        merge_key = None
        ws.merge_cells("A%d:%s%d" % (row, LAST, row))
        put("A%d" % row, rec, fill=F_SEC, font=Font(bold=True, size=10.5),
            align=Alignment(horizontal="left", vertical="center"))
        ws.row_dimensions[row].height = 20
        row += 1
        continue

    key = rec["A"]
    if key != merge_key:
        flush(row - 1)
        merge_start, merge_key = row, key
    mark = "● " if int(rec["B"]) in DRAWN else ""
    put("A%d" % row, mark + str(key), fill=F_KEY, font=Font(bold=True, size=11))
    for c in "BCDE":
        put("%s%d" % (c, row), rec[c], fill=F_KEY, font=Font(size=10))
    for c in "FGH":
        v = rec[c]
        put("%s%d" % (c, row), v if v != "" else "—", fill=F_REF,
            font=Font(size=10, color="404040" if v != "" else "9AA5B1"))
    for c in "IJKLM":
        v = rec.get(c, "")
        put("%s%d" % (c, row), v if v != "" else None,
            fill=(F_IN2 if c == "K" else F_IN),
            font=Font(size=11, bold=True, color=NAVY) if v != "" else Font(size=11))
    put("N%d" % row, rec["N"] or None, fill=F_IN,
        font=Font(size=9.5, color=NAVY) if rec["N"] else Font(size=9.5), align=LEFT)
    ws.row_dimensions[row].height = 23
    row += 1
flush(row - 1)

# ── その他（空欄）────────────────────────────────────────
ws.merge_cells("A%d:%s%d" % (row, LAST, row))
put("A%d" % row, "その他（表に行が無い組合せ・特殊なケース）　下型・材質・板厚も書いてください",
    fill=F_SEC, font=Font(bold=True, size=10.5),
    align=Alignment(horizontal="left", vertical="center"))
ws.row_dimensions[row].height = 20
row += 1
for _ in range(5):
    for c, _, _ in COLS:
        fill = F_IN
        if c in "FGH":
            fill = F_REF
        elif c == "K":
            fill = F_IN2
        put("%s%d" % (c, row), "—" if c in "FGH" else None, fill=fill,
            font=Font(size=10, color="9AA5B1"))
    ws.row_dimensions[row].height = 23
    row += 1

# ── 注記 ───────────────────────────────────────────────
row += 1
NOTES = [
    "※ ●＝断面の図面がある5型（V12・V25・V32・V40・V63）。この5型だけ「段差S 幾何」が出せます。他は図面をいただければ出します。",
    "※ 青字はこれまでに聞いた実績です。違っていたら直してください。",
    "※ 「フランジA 上限」は、その行の段差Sのときに取れる最大です。超えるとVの台に当たります。",
    "※ 備考の「S35の場合 A95」は、段差を35にすればフランジAを95まで伸ばせるという意味です。段差Sが変わると当たる場所が変わるためです。",
    "※ 記入後、このファイルをそのまま返してください。判断シートに入れます。",
]
for n in NOTES:
    ws.merge_cells("A%d:%s%d" % (row, LAST, row))
    put("A%d" % row, n, font=Font(size=9.5, color="404040"),
        align=Alignment(horizontal="left", vertical="center"), border=False)
    ws.row_dimensions[row].height = 14
    row += 1

# ── 印刷設定 ────────────────────────────────────────────
ws.page_setup.paperSize = 8                 # A3
ws.page_setup.orientation = "portrait"
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.print_area = "A1:%s%d" % (LAST, row - 1)
ws.print_title_rows = "%d:%d" % (HR1, HR2)
ws.print_options.horizontalCentered = True
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.35
ws.freeze_panes = "F%d" % (HR2 + 1)
ws.sheet_view.showGridLines = False

wb.save(DST)
print("書き出し:", DST, "／ 最終行", row - 1, "／ 幅合計", sum(w for _, _, w in COLS))
